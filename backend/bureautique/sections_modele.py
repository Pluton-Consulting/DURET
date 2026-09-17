"""
REMPLIR LES RUBRIQUES D'UN MODÈLE, AU LIEU DE LE VIDER (17/09).

Relevé de Noa (Duret, mémoire technique Domofrance) : « il a tendance soit à
en faire un autre sans reprendre la trame, soit à reprendre la trame mais à ne
modifier que le titre ». Les deux moteurs existaient, et aucun ne faisait ce
qu'on attend d'un mémoire :

  * la TRAME (`bureautique/trame.py`) ne sait que chercher-remplacer des
    chaînes : elle change la page de garde, jamais le corps ;
  * le COMPOSITEUR (`skills/documents_dossier.py`) sait rédiger, mais
    `preparer_modele` VIDAIT tout le corps du modèle pour n'en garder que les
    styles et les en-têtes — la garde, l'organigramme, la fiche signalétique,
    les moyens, tout ce qui fait le mémoire DE L'ENTREPRISE disparaissait.

Ce module tient le milieu, et il le tient DANS le document d'origine :

  1. `structure()` découpe le corps du modèle par ses titres de premier niveau
     (la GARDE = ce qui précède le premier titre) ;
  2. le plan de rédaction dit, rubrique par rubrique, si elle REPREND une
     rubrique du modèle telle quelle (`reprise_modele`) ou si elle est rédigée ;
  3. `assembler()` compose le corps final dans le modèle : la garde et les
     rubriques reprises NE BOUGENT PAS (images, tableaux, zones de texte,
     relations : rien n'est copié, donc rien ne se perd), et seules les
     rubriques rédigées viennent du rendu — du texte, des listes, des tableaux,
     des images simples, dans les styles du modèle.

Pourquoi copier du rendu VERS le modèle et pas l'inverse : ce qui vient du
rendu est simple et nous l'avons écrit ; ce qui vit dans le modèle d'un client
ne l'est pas (dessins ancrés, objets incorporés, numérotations). On ne déplace
que ce qu'on maîtrise.
"""
from __future__ import annotations

import io
import os
import re
from copy import deepcopy

_TITRE = re.compile(r"^(?:heading|titre)\s*(\d)\b", re.I)
EXTRAIT = 700


def _niveau(paragraphe) -> int | None:
    """Le niveau de titre d'un paragraphe, ou None. Le STYLE d'abord, le niveau
    de plan ensuite (des modèles titrent avec un style maison + outlineLvl)."""
    from docx.oxml.ns import qn
    try:
        nom = paragraphe.style.name if paragraphe.style is not None else ""
    except Exception:  # noqa: BLE001 — style abîmé : ce n'est pas un titre
        nom = ""
    m = _TITRE.match(nom or "")
    if m:
        return int(m.group(1))
    plan = paragraphe._p.find("./" + qn("w:pPr") + "/" + qn("w:outlineLvl"))
    if plan is not None:
        try:
            return int(plan.get(qn("w:val"))) + 1
        except (TypeError, ValueError):
            return None
    return None


def _niveau_par_style(paragraphe) -> int | None:
    try:
        nom = paragraphe.style.name if paragraphe.style is not None else ""
    except Exception:  # noqa: BLE001
        return None
    m = _TITRE.match(nom or "")
    return int(m.group(1)) if m else None


def _pages_du_sommaire(garde, titres: list[str]) -> dict[int, int]:
    """La page de DÉBUT de chaque rubrique, lue dans le sommaire du modèle (« ETUDES 8 »).
    Sert à savoir combien de pages pèsent les rubriques reprises : c'est ce qui reste
    qui se rédige. Rien de lisible → dictionnaire vide, l'appelant estime."""
    from docx.oxml.ns import qn
    norme = lambda t: re.sub(r"[^a-z0-9]+", "", _sans_accents(t))  # noqa: E731
    lignes = []
    for el in garde:
        for p in ([el] if el.tag == qn("w:p") else el.iter(qn("w:p"))):
            t = "".join((x.text or "") for x in p.iter(qn("w:t"))).strip()
            m = re.search(r"(\d{1,3})\s*$", t)
            if m:
                lignes.append((norme(t[:m.start()]), int(m.group(1))))
    pages = {}
    for i, titre in enumerate(titres):
        cle = norme(titre)
        if not cle:
            continue
        for texte, page in lignes:
            if texte.endswith(cle):
                pages[i] = page
                break
    return pages


def _sans_accents(t) -> str:
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", str(t or "").casefold())
                   if unicodedata.category(c) != "Mn")


def _decoupe(doc):
    """(garde, sections) — des listes d'ÉLÉMENTS vivants du corps."""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    corps = [el for el in doc.element.body if el.tag != qn("w:sectPr")]
    niveaux = {}
    for i, el in enumerate(corps):
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            n = _niveau(p)
            if n and p.text.strip():
                niveaux[i] = n
    if not niveaux:
        return corps, []
    haut = min(niveaux.values())
    debuts = [i for i, n in sorted(niveaux.items()) if n == haut]
    # UN PARAGRAPHE DE LISTE N'EST PAS UNE RUBRIQUE (17/09). Le modèle de Duret titre ses
    # rubriques en « Heading 1 », et trois lignes de liste (« Habillage de l'escalier »…)
    # portent un niveau de plan 1 par accident : elles coupaient « Principes de réalisation »
    # en quatre rubriques. Dès qu'un vrai STYLE de titre existe à ce niveau, lui seul découpe.
    styles = [i for i in debuts if _niveau_par_style(Paragraph(corps[i], doc)) == haut]
    if styles:
        debuts = styles
    garde = corps[:debuts[0]]
    sections = []
    for rang, debut in enumerate(debuts):
        fin = debuts[rang + 1] if rang + 1 < len(debuts) else len(corps)
        sections.append(corps[debut:fin])
    return garde, sections


def _texte(elements) -> str:
    from docx.oxml.ns import qn
    # Les fragments d'un MÊME paragraphe se collent (« PERSONNEL D » + « ’ENCADREMENT »
    # est un seul mot coupé par Word) ; les paragraphes, eux, se séparent.
    morceaux = []
    for el in elements:
        for p in ([el] if el.tag == qn("w:p") else el.iter(qn("w:p"))):
            morceaux.append("".join((t.text or "") for t in p.iter(qn("w:t"))))
    return re.sub(r"\s+", " ", " ".join(morceaux)).strip()


def structure(octets: bytes) -> dict:
    """Ce que le plan doit savoir du modèle : sa garde et ses rubriques.

    Rend {"garde": extrait, "sections": [{index, titre, extrait, mots, images,
    tableaux}]} — vide de sections si le modèle n'a aucun titre : l'appelant
    garde alors l'ancien comportement, il ne devine pas un découpage."""
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(octets))
    garde, sections = _decoupe(doc)
    fiches = []
    for i, elements in enumerate(sections):
        titre = _texte(elements[:1])
        corps = _texte(elements[1:])
        fiches.append({
            "index": i, "titre": titre, "extrait": corps[:EXTRAIT], "mots": len(corps.split()),
            "images": sum(len(el.findall(".//" + qn("w:drawing"))) + len(el.findall(".//" + qn("w:pict")))
                          for el in elements),
            "tableaux": sum(1 for el in elements if el.tag == qn("w:tbl")),
        })
    # LE POIDS EN PAGES de chaque rubrique : lu au sommaire quand il existe (l'écart entre
    # deux débuts), estimé sinon (mots, images, tableaux). La dernière rubrique n'a pas de
    # suivante : toujours estimée.
    debuts = _pages_du_sommaire(garde, [f["titre"] for f in fiches])
    for i, f in enumerate(fiches):
        estime = round(max(0.3, f["mots"] / 320 + f["images"] * 0.35 + f["tableaux"] * 0.3), 1)
        if i in debuts and i + 1 in debuts and debuts[i + 1] >= debuts[i]:
            f["pages"] = max(0.3, float(debuts[i + 1] - debuts[i])) if debuts[i + 1] > debuts[i] else min(estime, 0.6)
        else:
            f["pages"] = estime
    pages_garde = max(1, debuts[0] - 1) if 0 in debuts else 2
    return {"garde": _texte(garde)[:EXTRAIT], "sections": fiches, "pages_garde": pages_garde}


def textes_des_rubriques(octets: bytes, index: list[int], plafond: int = 12000) -> dict[int, str]:
    """Le TEXTE ENTIER des rubriques demandées, paragraphe par paragraphe : le rédacteur
    d'une rubrique de chantier part de la façon de faire écrite par l'entreprise, il ne
    la réinvente pas. `structure()` n'en rend qu'un extrait de 700 caractères."""
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(octets))
    _, sections = _decoupe(doc)
    textes = {}
    for i in index:
        if not 0 <= i < len(sections):
            continue
        lignes = []
        for el in sections[i][1:]:
            for p in ([el] if el.tag == qn("w:p") else el.iter(qn("w:p"))):
                t = "".join((x.text or "") for x in p.iter(qn("w:t"))).strip()
                if t:
                    lignes.append(t)
        textes[i] = "\n".join(lignes)[:plafond]
    return textes


def _groupes_du_rendu(rendu, titres: list[str]) -> dict[int, list]:
    """Les éléments du rendu, rangés par rubrique du plan. Une rubrique
    commence au paragraphe dont le texte EST son titre ; ce qui précède la
    première (garde et sommaire fabriqués par le rendu) est laissé de côté :
    le modèle a sa propre garde."""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    norme = lambda s: re.sub(r"\s+", " ", str(s or "")).strip().casefold()  # noqa: E731
    attendus = [norme(t) for t in titres]
    groupes: dict[int, list] = {}
    courant, suivant = None, 0
    for el in rendu.element.body:
        if el.tag == qn("w:sectPr"):
            continue
        if el.tag == qn("w:p") and suivant < len(attendus):
            if norme(Paragraph(el, rendu).text) == attendus[suivant]:
                courant, suivant = suivant, suivant + 1
                groupes[courant] = []
        if courant is not None:
            groupes[courant].append(el)
    return groupes


def _greffer(element, rendu, modele):
    """Une copie de `element` (venu du rendu) utilisable DANS le modèle : ses
    images sont réinscrites dans le paquet du modèle. Toute autre relation
    (objet incorporé…) n'existe pas dans ce que le rendu fabrique."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    copie = deepcopy(element)
    for blip in copie.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        partie = rendu.part.related_parts.get(rid) if rid else None
        if partie is None:
            continue
        nouveau, _ = modele.part.get_or_add_image(io.BytesIO(partie.blob))
        blip.set(qn("r:embed"), nouveau)
    for lien in copie.iter(qn("w:hyperlink")):
        rid = lien.get(qn("r:id"))
        rel = rendu.part.rels.get(rid) if rid else None
        if rel is not None and rel.is_external:
            lien.set(qn("r:id"), modele.part.relate_to(rel.target_ref, RT.HYPERLINK, is_external=True))
        elif rid:
            del lien.attrib[qn("r:id")]
    return copie


def assembler(original: bytes, rendu_chemin: str, titres: list[str], reprises: dict) -> dict:
    """Compose le document final DANS le modèle et l'écrit à la place du rendu.

    `titres` : les rubriques du plan, dans l'ordre. `reprises` : {titre du plan
    → index de la rubrique du modèle gardée telle quelle}. Rend un compte
    rendu ; ne touche à rien si le modèle n'a pas de rubriques."""
    from docx import Document
    from docx.oxml.ns import qn
    modele = Document(io.BytesIO(original))
    garde, sections = _decoupe(modele)
    if not sections:
        return {"assemble": False, "raison": "le modèle n'a aucun titre : sa présentation seule est reprise"}
    rendu = Document(rendu_chemin)
    groupes = _groupes_du_rendu(rendu, titres)
    norme = lambda s: re.sub(r"\s+", " ", str(s or "")).strip().casefold()  # noqa: E731
    reprises = {norme(k): v for k, v in (reprises or {}).items()}

    corps = modele.element.body
    final_sect = corps.find("./" + qn("w:sectPr"))
    nouveaux = list(garde)
    reprises_faites, redigees, manquantes = [], [], []
    for rang, titre in enumerate(titres):
        index = reprises.get(norme(titre))
        if isinstance(index, int) and 0 <= index < len(sections):
            nouveaux.extend(sections[index])
            reprises_faites.append(index)
        elif rang in groupes:
            nouveaux.extend(_greffer(el, rendu, modele) for el in groupes[rang])
            redigees.append(rang)
        else:
            manquantes.append(titre)
    # Une rubrique du modèle écartée peut porter un SAUT DE SECTION (orientation,
    # en-tête différent) : on garde le saut, pas son texte — sinon les pages qui
    # suivent héritent de la mauvaise mise en page.
    gardes = {id(el) for el in nouveaux}
    for index, elements in enumerate(sections):
        if index in reprises_faites:
            continue
        for el in elements:
            if id(el) in gardes or el.tag != qn("w:p"):
                continue
            if el.find("./" + qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
                for enfant in list(el):
                    if enfant.tag != qn("w:pPr"):
                        el.remove(enfant)
                nouveaux.append(el)
    for el in list(corps):
        if el is not final_sect:
            corps.remove(el)
    for el in nouveaux:
        if final_sect is not None:
            final_sect.addprevious(el)
        else:
            corps.append(el)
    # LES IMAGES DE L'ANCIEN CHANTIER NE RESTENT PAS CACHÉES DANS LE FICHIER :
    # une rubrique écartée laisse sa relation, donc son image, dans le paquet —
    # invisible à l'écran, lisible par qui ouvre le .docx comme une archive.
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from lxml import etree
    encore = etree.tostring(corps).decode("utf-8", "ignore")
    retirees = 0
    for rid, rel in list(modele.part.rels.items()):
        if rel.reltype == RT.IMAGE and f'"{rid}"' not in encore:
            del modele.part.rels[rid]
            retirees += 1
    tmp = f"{rendu_chemin}.{os.getpid()}.assemble.tmp"
    # LA TABLE DES MATIÈRES DU MODÈLE SE RECALCULE À L'OUVERTURE : ses lignes sont du
    # texte figé (celles de l'ancien document) ; Word les refait dès qu'on l'y invite.
    try:
        reglages = modele.settings.element
        if reglages.find(qn("w:updateFields")) is None:
            champ = reglages.makeelement(qn("w:updateFields"), {qn("w:val"): "true"})
            reglages.append(champ)
    except Exception:  # noqa: BLE001 — un confort, jamais un motif d'échec du rendu
        pass
    modele.save(tmp)
    Document(tmp)            # le fichier produit se ROUVRE, ou il n'est pas publié
    os.replace(tmp, rendu_chemin)
    return {"assemble": True, "garde_conservee": bool(garde), "rubriques_reprises": len(reprises_faites),
            "rubriques_redigees": len(redigees), "rubriques_sans_contenu": manquantes,
            "images_ecartees": retirees}


def copier_entete(document: bytes, source: bytes) -> tuple[bytes, dict]:
    """L'EN-TÊTE ET LE PIED D'UN WORD, POSÉS SUR UN AUTRE (17/09).

    « Mets en en-tête l'en-tête de Revêtements Duret Sols » : l'en-tête de la
    maison est un DOCUMENT WORD (logo ancré + mentions légales en pied), pas
    une image. On recopie les paragraphes de l'en-tête et du pied de `source`
    dans chaque section de `document`, images comprises. Le corps ne bouge pas."""
    from docx import Document
    from docx.oxml.ns import qn
    cible = Document(io.BytesIO(document))
    origine = Document(io.BytesIO(source))
    modele = origine.sections[0]
    poses = {"entete": 0, "pied": 0, "images": 0}

    def _copier(depuis, vers, cle):
        if depuis.is_linked_to_previous and not list(depuis._element):
            return
        vers.is_linked_to_previous = False
        for enfant in list(vers._element):
            vers._element.remove(enfant)
        for enfant in depuis._element:
            copie = deepcopy(enfant)
            for blip in copie.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                partie = depuis.part.related_parts.get(rid) if rid else None
                if partie is not None:
                    nouveau, _ = vers.part.get_or_add_image(io.BytesIO(partie.blob))
                    blip.set(qn("r:embed"), nouveau)
                    poses["images"] += 1
            for image in copie.iter("{urn:schemas-microsoft-com:vml}imagedata"):
                rid = image.get(qn("r:id"))
                partie = depuis.part.related_parts.get(rid) if rid else None
                if partie is not None:
                    nouveau, _ = vers.part.get_or_add_image(io.BytesIO(partie.blob))
                    image.set(qn("r:id"), nouveau)
                    poses["images"] += 1
            vers._element.append(copie)
        poses[cle] += 1

    for section in cible.sections:
        _copier(modele.header, section.header, "entete")
        _copier(modele.footer, section.footer, "pied")
        if modele.different_first_page_header_footer:
            section.different_first_page_header_footer = True
            _copier(modele.first_page_header, section.first_page_header, "entete")
            _copier(modele.first_page_footer, section.first_page_footer, "pied")
        # Les marges d'en-tête suivent : un logo haut de 3 cm dans une marge
        # de 1,25 cm recouvre le texte.
        for attribut in ("top_margin", "bottom_margin", "header_distance", "footer_distance"):
            valeur = getattr(modele, attribut)
            if valeur is not None:
                setattr(section, attribut, valeur)
    sortie = io.BytesIO()
    cible.save(sortie)
    Document(io.BytesIO(sortie.getvalue()))
    return sortie.getvalue(), poses



# ─── LA GARDE D'UN MODÈLE, ÉCRITE EN PARAGRAPHES SÉPARÉS ────────────────────
#
# 17/09, modèle réel de Duret : la garde porte « Projet : » / « Réaménagement… »
# / « a le TAILLAN MEDOC » / « Maître d’ouvrage : » / « BORDEAUX METROPOLE » /
# l'adresse… chacun dans SON paragraphe (et en zones de texte). Le plan, lui,
# donne « Projet : Réaménagement … TAILLAN MEDOC » → « Projet : Construction … » :
# ce texte d'un seul tenant n'existe nulle part, 0 remplacement sur 4, la garde
# de l'ancien chantier restait — et le contrôle final le re-signalait à chaque
# tour, sans que rien puisse le corriger.

def _norme_garde(t: str) -> str:
    t = str(t or "").replace("\u00a0", " ").replace("’", "'").casefold()
    return " ".join(t.split())


def _texte_p(p, qn) -> str:
    return "".join(x.text or "" for x in p.iter(qn("w:t")))


def _ecrire_p(p, texte: str, qn) -> None:
    noeuds = list(p.iter(qn("w:t")))
    if not noeuds:
        return
    noeuds[0].text = texte
    noeuds[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for n in noeuds[1:]:
        n.text = ""


def actualiser_garde(octets: bytes, table: dict) -> tuple[bytes, int]:
    """Applique à la GARDE les remplacements « Libellé : ancienne valeur » →
    « Libellé : nouvelle valeur » quand libellé et valeur vivent dans des
    paragraphes séparés. Rend (octets, nombre de libellés actualisés). Ne touche
    à rien après le premier titre."""
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(octets))
    elements, _sections = _decoupe(doc)
    garde = [p for e in elements for p in ([e] if e.tag == qn("w:p") else []) + [x for x in e.iter(qn("w:p")) if x is not e]]
    # un paragraphe qui en contient d'autres (zone de texte) n'est pas une ligne de texte
    garde = [p for p in garde if not any(True for x in p.iter(qn("w:p")) if x is not p)]
    libelle_de = lambda t: _norme_garde(t).split(":", 1)[0].strip() if ":" in t else ""
    faits = 0
    for ancien, nouveau in (table or {}).items():
        lib = libelle_de(str(ancien))
        if not lib or libelle_de(str(nouveau)) != lib or len(lib) > 40:
            continue
        valeur = str(nouveau).split(":", 1)[1].strip()
        fait = False
        for i, p in enumerate(garde):
            t = _texte_p(p, qn)
            if libelle_de(t) != lib:
                continue
            reste = t.split(":", 1)[1].strip()
            if reste:                                   # « Date : le 24/07/2026 » : tout dans le même paragraphe
                _ecrire_p(p, t.split(":", 1)[0] + ": " + valeur, qn)
                fait = True
                continue
            pose = False                                # la valeur suit, en un ou plusieurs paragraphes
            for q in garde[i + 1:i + 8]:
                tq = _texte_p(q, qn)
                if not tq.strip():
                    continue
                if libelle_de(tq) and len(libelle_de(tq)) <= 40:
                    break
                _ecrire_p(q, valeur if not pose else "", qn)
                pose = True
            fait = fait or pose
        faits += int(fait)
    if not faits:
        return octets, 0
    sortie = io.BytesIO()
    doc.save(sortie)
    Document(io.BytesIO(sortie.getvalue()))
    return sortie.getvalue(), faits
