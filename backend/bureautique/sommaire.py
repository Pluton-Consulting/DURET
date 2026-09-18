"""
LE SOMMAIRE D'UN MODÈLE REMPLI DIT CE QUE LE DOCUMENT CONTIENT (18/09, banc Duret, Q19).

« Refais le mémoire uniquement pour le lot 12 » : le corps ne parlait plus que des
sols souples, et l'aperçu du chat montrait en page 2 un sommaire listant
« 1) Carrelages et faïence », « f) Pose du carrelage au sol »… — les lignes FIGÉES
du sommaire de la trame, celles de l'ancien chantier. Un sommaire Word est un
champ : ses lignes sont du texte mis en cache, que seul Word recalcule, et
seulement à l'ouverture (`updateFields`). L'aperçu du chat, le PDF et LibreOffice
affichent le cache.

Ce module refait le cache à partir du document lui-même, mécaniquement :
  1. les TITRES réels du corps, dans l'ordre, aux niveaux que le champ couvre
     (`TOC \\o "1-3"`) ;
  2. un rendu PDF (LibreOffice, déjà utilisé pour compter les pages) donne pour
     chacun sa PAGE et son NUMÉRO affiché (« I. », « 1) », « a) » — la
     numérotation automatique n'existe qu'au rendu) ;
  3. les lignes du sommaire sont réécrites dans les styles de la trame (TM1,
     TM2… avec leurs taquets), le champ restant un champ : Word le recalcule
     toujours à l'ouverture.
Deux passes au plus : réécrire le sommaire peut déplacer les pages qui suivent.
Tout échec laisse le fichier tel quel : un sommaire périmé ne vaut pas un
document perdu.
"""
from __future__ import annotations

import logging
import os
import re
from copy import deepcopy

logger = logging.getLogger("infra.sommaire")

_NIVEAUX_TOC = re.compile(r'\\o\s*"(\d)\s*-\s*(\d)"')
_STYLE_TITRE = re.compile(r"^(?:heading|titre)\s*(\d)\b", re.I)
_ETIQUETTE = re.compile(r"^(?:[ivxlcdm]{1,6}\.|\d{1,3}[.)]|\d{1,2}(?:\.\d{1,2}){1,3}\.?|[a-z][.)])$", re.I)
_LIGNE_DE_SOMMAIRE = re.compile(r"(?:\.{4,}|…{2,})\s*\d{1,4}\s*$")
_NUMERO_EN_TETE = re.compile(r"^\s*((?:[ivxlcdm]{1,6}\.|\d{1,3}[.)]|\d{1,2}(?:\.\d{1,2}){1,3}\.?|[a-z][.)]))\s+", re.I)


def _norme(t: str) -> str:
    return " ".join(str(t or "").replace("’", "'").split()).casefold()


def _texte(p, qn) -> str:
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def _styles(doc) -> dict:
    """identifiant de style → (nom, niveau de plan éventuel)."""
    from docx.oxml.ns import qn
    table = {}
    for st in doc.styles.element.iter(qn("w:style")):
        ident = st.get(qn("w:styleId"))
        nom_el = st.find(qn("w:name"))
        nom = nom_el.get(qn("w:val")) if nom_el is not None else ""
        lvl = st.find("./" + qn("w:pPr") + "/" + qn("w:outlineLvl"))
        table[ident] = (nom or "", int(lvl.get(qn("w:val"))) + 1 if lvl is not None and str(lvl.get(qn("w:val"))).isdigit() else None)
    return table


def _style_id(p, qn):
    s = p.find("./" + qn("w:pPr") + "/" + qn("w:pStyle"))
    return s.get(qn("w:val")) if s is not None else None


def _niveau(p, styles, qn):
    ident = _style_id(p, qn)
    nom, lvl = styles.get(ident, ("", None))
    m = _STYLE_TITRE.match(nom or "") or _STYLE_TITRE.match(ident or "")
    if m:
        return int(m.group(1))
    propre = p.find("./" + qn("w:pPr") + "/" + qn("w:outlineLvl"))
    if propre is not None and str(propre.get(qn("w:val"))).isdigit():
        return int(propre.get(qn("w:val"))) + 1
    return lvl


def _champ_du_sommaire(corps, qn):
    """(paragraphe qui ouvre le champ TOC, niveaux couverts) ou (None, None)."""
    for p in corps.iter(qn("w:p")):
        for instr in p.iter(qn("w:instrText")):
            if re.match(r"\s*TOC\b", instr.text or ""):
                m = _NIVEAUX_TOC.search(instr.text or "")
                return p, (int(m.group(1)), int(m.group(2))) if m else (1, 3)
    return None, None


def _fin_du_champ(debut, qn):
    """Le paragraphe où le champ du sommaire se FERME : on compte les ouvertures et
    fermetures de champs (les PAGEREF des lignes sont des champs imbriqués)."""
    profondeur = 0
    p = debut
    while p is not None:
        if p.tag == qn("w:p"):
            for fc in p.iter(qn("w:fldChar")):
                t = fc.get(qn("w:fldCharType"))
                if t == "begin":
                    profondeur += 1
                elif t == "end":
                    profondeur -= 1
                    if profondeur == 0:
                        return p
        p = p.getnext()
    return None


def _pages_et_numeros(pdf: bytes, titres: list[str]) -> list[tuple]:
    """Pour chaque titre (dans l'ordre) : (page, numéro affiché) ou (None, None)."""
    import fitz
    sortie = []
    with fitz.open(stream=pdf, filetype="pdf") as d:
        pages = [[l for l in d[i].get_text("text").splitlines() if l.strip()] for i in range(d.page_count)]
    # Le corps commence après le sommaire : les pages qui portent des lignes « titre ….. 12 ».
    debut = 0
    for i, lignes in enumerate(pages[:12]):
        if sum(1 for l in lignes if _LIGNE_DE_SOMMAIRE.search(l)) >= 3:
            debut = i + 1
    page = debut
    for titre in titres:
        cible = _norme(titre)
        trouve = (None, None)
        for i in range(page, len(pages)):
            lignes = pages[i]
            for k, ligne in enumerate(lignes):
                if _LIGNE_DE_SOMMAIRE.search(ligne) or not ligne.strip():
                    continue
                # Le numéro peut précéder le titre sur la MÊME ligne (« 7) Contrôle de la
                # planéité ») ou sur la ligne d'avant (« I. » seul) : les deux se lisent.
                m = _NUMERO_EN_TETE.match(ligne)
                numero = m.group(1) if m else None
                n = _norme(ligne[m.end():] if m else ligne)
                if not n or not (n == cible or (len(cible) > 20 and n.startswith(cible[:50])) or (len(n) >= 15 and cible.startswith(n))):
                    continue
                if numero is None and k and _ETIQUETTE.match(lignes[k - 1].strip()):
                    numero = lignes[k - 1].strip()
                trouve = (i + 1, numero)
                break
            if trouve[0]:
                page = trouve[0] - 1
                break
        sortie.append(trouve)
    return sortie


def _reecrire(doc, entrees: list[tuple], qn) -> bool:
    """Réécrit les lignes du sommaire : entrees = [(niveau, titre, page, numéro)]."""
    from docx.oxml import OxmlElement
    corps = doc.element.body
    debut, _ = _champ_du_sommaire(corps, qn)
    if debut is None:
        return False
    parent = debut.getparent()
    prefixe = re.sub(r"\d+$", "", _style_id(debut, qn) or "")
    fin = _fin_du_champ(debut, qn)
    if fin is None or fin is debut or not prefixe:
        return False
    anciennes = [debut]
    suivant = debut.getnext()
    while suivant is not None and suivant is not fin and suivant.tag == qn("w:p") and re.fullmatch(re.escape(prefixe) + r"\d", _style_id(suivant, qn) or ""):
        anciennes.append(suivant)
        suivant = suivant.getnext()
    if suivant is fin and re.fullmatch(re.escape(prefixe) + r"\d", _style_id(fin, qn) or ""):
        anciennes.append(fin)          # la dernière ligne porte elle-même la fermeture du champ
        fin_dans_la_derniere = True
    elif suivant is fin:
        fin_dans_la_derniere = False   # la fermeture est dans son propre paragraphe : il reste
    else:
        return False                   # autre chose entre les lignes et la fermeture : on ne touche pas
    # Gabarits par niveau : la mise en forme (style, taquets) et le texte (police) de la trame.
    gabarits = {}
    for p in anciennes:
        m = re.search(r"(\d)$", _style_id(p, qn) or "")
        if not m or int(m.group(1)) in gabarits:
            continue
        ppr = p.find(qn("w:pPr"))
        rpr = None
        for r in p.iter(qn("w:r")):
            if r.find(qn("w:t")) is not None and (r.find(qn("w:t")).text or "").strip():
                rpr = r.find(qn("w:rPr"))
                break
        gabarits[int(m.group(1))] = (ppr, rpr)
    if not gabarits:
        return False
    # Les trois pièces qui OUVRENT le champ restent (début, instruction, séparateur).
    ouverture = []
    for r in debut.findall(qn("w:r")):
        ouverture.append(r)
        fc = r.find(qn("w:fldChar"))
        if fc is not None and fc.get(qn("w:fldCharType")) == "separate":
            break

    def run(texte=None, tab=False, rpr=None):
        r = OxmlElement("w:r")
        if rpr is not None:
            r.append(deepcopy(rpr))
        if tab:
            r.append(OxmlElement("w:tab"))
        else:
            t = OxmlElement("w:t")
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            t.text = texte
            r.append(t)
        return r

    neufs = []
    for rang, (niveau, titre, page, numero) in enumerate(entrees):
        cle = niveau if niveau in gabarits else min(gabarits, key=lambda k: abs(k - niveau))
        ppr, rpr = gabarits[cle]
        p = OxmlElement("w:p")
        if ppr is not None:
            p.append(deepcopy(ppr))
        if rang == 0:
            for r in ouverture:
                p.append(deepcopy(r))
        if numero:
            # Un numéro plus large que le premier taquet (« XVIII. ») enverrait le titre
            # derrière les points de conduite : une espace le suit alors, pas une tabulation.
            p.append(run(numero + (" " if len(numero) >= 5 else ""), rpr=rpr))
            if len(numero) < 5:
                p.append(run(tab=True, rpr=rpr))
        p.append(run(titre, rpr=rpr))
        p.append(run(tab=True, rpr=rpr))
        p.append(run(str(page) if page else "", rpr=rpr))
        neufs.append(p)
    if fin_dans_la_derniere:
        fin = OxmlElement("w:r")
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), "end")
        fin.append(fc)
        neufs[-1].append(fin)
    position = parent.index(anciennes[0])
    for p in anciennes:
        parent.remove(p)
    for decalage, p in enumerate(neufs):
        parent.insert(position + decalage, p)
    return True


def _titres(doc, niveaux, qn):
    """Les titres du corps APRÈS le sommaire, dans l'ordre, aux niveaux qu'il couvre."""
    corps = doc.element.body
    debut, _ = _champ_du_sommaire(corps, qn)
    fin = _fin_du_champ(debut, qn) if debut is not None else None
    styles = _styles(doc)
    passe_le_sommaire = debut is None
    vus = []
    for p in corps.iter(qn("w:p")):
        if not passe_le_sommaire:
            if p is fin:
                passe_le_sommaire = True
            continue
        niveau = _niveau(p, styles, qn)
        texte = " ".join(_texte(p, qn).split())
        if niveau and niveaux[0] <= niveau <= niveaux[1] and texte:
            vus.append((niveau, texte))
    return vus


def actualiser(chemin: str, convertir=None) -> dict:
    """Refait les lignes du sommaire d'un .docx d'après ses titres et leur page rendue."""
    from docx import Document
    from docx.oxml.ns import qn
    try:
        doc = Document(chemin)
        debut, niveaux = _champ_du_sommaire(doc.element.body, qn)
        if debut is None:
            return {"sommaire": "absent"}
        titres = _titres(doc, niveaux, qn)
        if not titres:
            return {"sommaire": "sans titres"}
        if convertir is None:
            from bureautique.document_modele import convertir_pdf as convertir
        pdf = convertir(chemin)
        reperes = _pages_et_numeros(pdf, [t for _, t in titres])
        entrees = [(n, t, pg, num) for (n, t), (pg, num) in zip(titres, reperes)]
        tmp = f"{chemin}.{os.getpid()}.sommaire.tmp"
        for passe in (1, 2):
            if not _reecrire(doc, entrees, qn):
                return {"sommaire": "structure non reconnue"}
            doc.save(tmp)
            Document(tmp)                      # le fichier se rouvre, ou il n'est pas publié
            if passe == 2:
                break
            reperes2 = _pages_et_numeros(convertir(tmp), [t for _, t in titres])
            if [p for p, _ in reperes2] == [p for p, _ in reperes]:
                break
            entrees = [(n, t, pg, num or ancien) for (n, t), (pg, num), (_, ancien) in zip(titres, reperes2, reperes)]
            doc = Document(tmp)
        os.replace(tmp, chemin)
        sans_page = sum(1 for e in entrees if not e[2])
        return {"sommaire": "actualisé", "lignes": len(entrees), "sans_page": sans_page}
    except Exception as e:  # noqa: BLE001 — un sommaire périmé ne vaut pas un document perdu
        logger.warning("Sommaire non actualisé (%s) : %s", type(e).__name__, str(e)[:160])
        try:
            if "tmp" in locals() and os.path.exists(tmp):
                os.remove(tmp)
        except OSError:
            pass
        return {"sommaire": "non actualisé", "raison": type(e).__name__}
