"""
Banc « UN CLASSEUR QUI SE CALCULE » (17/09). Rejoué sur le code livré : dans un Excel écrit par
l'assistant, toutes les cellules étaient du TEXTE (« 412,50 », « 38,00 € », et même 412.5 donné en
nombre). Le typage est PRUDENT : une colonne ne devient numérique que si toutes ses cellules
remplies sont des nombres sans ambiguïté, de même unité, sous un en-tête qui n'est pas un
identifiant ; et une cellule qui commence par « = » n'est une formule que si elle est sûre.
Vrai rendu openpyxl, classeur ROUVERT. Sans base ni réseau.

    python backend/scripts/test_excel_calculable.py backend
"""
import os, sys, tempfile
from pathlib import Path

BACKEND = Path(sys.argv[1] if len(sys.argv) > 1 else "backend").resolve()
sys.path.insert(0, str(BACKEND))
for k, v in {"DATABASE_URL": "postgresql://x:x@localhost/x", "JWT_SECRET_KEY": "x" * 40, "RESEND_API_KEY": "x"}.items():
    os.environ.setdefault(k, v)
ECHECS = []


def verifier(nom, condition, detail=""):
    print(("  ✓ " if condition else "  ✗ ") + nom + (f"  → {detail}" if detail and not condition else ""))
    if not condition:
        ECHECS.append(nom)


try:
    import openpyxl
except ImportError:
    print("openpyxl absent de ce poste : banc sauté."); sys.exit(0)
from bureautique.modele import normaliser_element, normaliser_entete, deplier_feuilles
from bureautique import rendu


def classeur(blocs):
    elements = [e for e in (normaliser_element(x) for x in deplier_feuilles(blocs)) if e]
    chemin = os.path.join(tempfile.mkdtemp(), "essai.xlsx")
    rendu._xlsx(normaliser_entete({"titre": "Essai", "format": "xlsx"}), elements, chemin)
    return openpyxl.load_workbook(chemin)


print("1. Les nombres")
wb = classeur([{"bloc": "feuille", "nom": "Métrés", "total": True,
                "entetes": ["Poste", "Qté", "PU", "Total", "Téléphone", "Code postal", "N° affaire", "Avancement"],
                "lignes": [["Carrelage 60x60", "412,50", "38,00 €", "=B2*C2", "0612345678", "33260", "2026-71", "35 %"],
                           ["PVC logements", "1 230,10", "22,5 €", "=B3*C3", "05 56 00 00 00", "01000", "2026-72", "100 %"]]}])
ws = wb["Métrés"]
verifier("une quantité écrite à la française devient un NOMBRE (412,50 → 412.5 ; 1 230,10 → 1230.1)", ws["B2"].value == 412.5 and ws["B3"].value == 1230.1, str((ws["B2"].value, ws["B3"].value)))
verifier("un prix garde son unité à l'AFFICHAGE, pas dans la valeur", ws["C2"].value == 38.0 and "€" in ws["C2"].number_format)
verifier("une formule arithmétique reste une formule", ws["D2"].value == "=B2*C2" and ws["D2"].data_type == "f")
verifier("un téléphone, un code postal, un numéro d'affaire restent du TEXTE", ws["E2"].value == "0612345678" and ws["F3"].value == "01000" and ws["G2"].value == "2026-71")
verifier("un pourcentage se calcule (35 % → 0,35 affiché 35 %)", abs(ws["H2"].value - 0.35) < 1e-9 and ws["H2"].number_format == "0.00%")
verifier("`total: true` pose une ligne de TOTAL par formule, qui se recalcule", ws["A4"].value == "Total" and ws["B4"].value == "=SUM(B2:B3)" and ws["C4"].value == "=SUM(C2:C3)", str([c.value for c in ws[4]]))

print("2. La prudence")
wb = classeur([{"bloc": "feuille", "nom": "Mixte", "entetes": ["Poste", "Qté", "Note"],
                "lignes": [["a", "12", "3 passes"], ["b", "environ 14", "2"], ["c", "", "1"]]}])
ws = wb["Mixte"]
verifier("UNE cellule ambiguë (« environ 14 ») et toute la colonne reste du texte", ws["B2"].value == "12" and ws["B3"].value == "environ 14")
verifier("des unités différentes dans une colonne : rien n'est converti", classeur([{"bloc": "feuille", "nom": "U", "entetes": ["x", "Mesure"], "lignes": [["a", "12 m²"], ["b", "14 ml"]]}])["U"]["B2"].value == "12 m²")
verifier("une même unité partout : convertie, unité à l'affichage", classeur([{"bloc": "feuille", "nom": "U", "entetes": ["x", "Surface"], "lignes": [["a", "12,4 m²"], ["b", "14 m2"]]}])["U"]["B2"].value == 12.4)
verifier("de longs entiers sans en-tête de quantité sont des identifiants", classeur([{"bloc": "feuille", "nom": "I", "entetes": ["x", "Dossier"], "lignes": [["a", "797427150"], ["b", "797427151"]]}])["I"]["B2"].value == "797427150")

print("3. Aucune formule qu'on n'a pas écrite")
wb = classeur([{"bloc": "feuille", "nom": "Mails", "entetes": ["Objet", "Note"],
                "lignes": [['=HYPERLINK("http://exemple.invalid","clic")', "+33 5 56"], ["=cmd|' /C calc'!A0", "@rappel"], ["=SOMME(B2;B3)", "ok"]]}])
ws = wb["Mails"]
verifier("un objet de mail qui commence par « = » reste du TEXTE", ws["A2"].data_type == "s" and ws["A3"].data_type == "s", str((ws["A2"].data_type, ws["A3"].data_type)))
verifier("« + » et « @ » en tête de cellule restent du texte", ws["B2"].data_type == "s" and ws["B3"].data_type == "s")
verifier("une fonction de total écrite en français est comprise (SOMME ; → SUM ,)", ws["A4"].value == "=SUM(B2,B3)" and ws["A4"].data_type == "f", str(ws["A4"].value))

print("4. Le quantitatif garde ses colonnes déclarées, et les formes naturelles font leurs onglets")
wb = classeur([{"bloc": "feuille", "nom": "Détail", "colonnes_numeriques": [1], "entetes": ["Local", "Quantité"], "lignes": [["Séjour", "25.9"], ["Chambre", "13.4"]]}])
verifier("`colonnes_numeriques` (quantitatif) donne toujours des nombres", wb["Détail"]["B2"].value == 25.9)
wb = classeur([{"type": "feuille", "nom": "Lot 11", "contenu": [{"type": "tableau", "colonnes": ["Poste", "Qté"], "lignes": [["Carrelage", "412,5"]]}]},
               {"bloc": "feuille", "titre": "Lot 12"}, {"bloc": "tableau", "entetes": ["Poste", "Qté"], "lignes": [["PVC", "1 230,1"]]}])
verifier("feuille qui PORTE ses blocs + feuille-séparateur : deux onglets, nombres compris", wb.sheetnames == ["Lot 11", "Lot 12"] and wb["Lot 11"]["B2"].value == 412.5 and wb["Lot 12"]["B2"].value == 1230.1, str(wb.sheetnames))

print(("✗ %d échec(s) : %s" % (len(ECHECS), ", ".join(ECHECS))) if ECHECS else "✓ 0 échec")
sys.exit(1 if ECHECS else 0)
