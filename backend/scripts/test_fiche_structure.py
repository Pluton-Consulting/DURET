"""
Banc de la fiche de lecture du 21/09 (prompt 1 du cahier « Dix-neuf prompts » de Duret).

« Fiche de lecture du DCE Mouriscot » est sortie :
  · avec la page de garde d'un AUTRE projet (le Word « MEMOIRE Complet Vierge.docx » du dossier
    était imposé comme trame à n'importe quel document) ;
  · sans le RC ni le CCAP, des « .doc » que le rédacteur écartait — date limite et pénalités
    déclarées absentes ;
  · avec la synthèse et le contrôle écrits deux fois (la première rubrique avait tout rédigé) ;
  · avec un tableau à quatre colonnes égales, la colonne des citations étirée sur une page.

Éprouvé ici, les fonctions EXÉCUTÉES :
  · `_imposer_modele` : un Word maison n'est une TRAME que si la demande le désigne ; sinon
    présentation seule ;
  · `largeurs_colonnes` et `_tableau_docx` (python-docx) : parts proportionnées, marges de
    cellules, citations d'une cellule à la ligne ;
  · `lecture_integrale.lire` sur un vrai « .doc » (LibreOffice), là où il est installé.

Usage : python backend/scripts/test_fiche_structure.py [backend]
"""
import ast
import asyncio
import importlib.util
import io
import logging
import pathlib
import re
import sys
import types

BACKEND = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "backend").resolve()
sys.path.insert(0, str(BACKEND))
echecs = []


def verifier(nom, cond, detail=""):
    print(f"  {'✓' if cond else '✗'} {nom}" + (f"  → {detail}" if detail and not cond else ""))
    if not cond:
        echecs.append(nom)


def fonctions(chemin, noms, espace):
    """Extrait des fonctions et des affectations de module d'un fichier lourd, et les exécute."""
    source = (BACKEND / chemin).read_text(encoding="utf-8")
    arbre = ast.parse(source)
    garder = [n for n in arbre.body
              if (isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name in noms)
              or (isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id in noms for t in n.targets))]
    exec(compile(ast.Module(body=garder, type_ignores=[]), chemin, "exec"), espace)
    return espace


print(f"\n═══ LA FICHE DE LECTURE — {BACKEND.parent}\n")
composeur = (BACKEND / "skills" / "documents_dossier.py").read_text(encoding="utf-8")

# ── 1. Le Word maison : trame seulement si la demande le désigne ──────────────
class _Dossiers:
    def __init__(self, sources):
        self._s = sources

    def sources(self, uid, fil, ids=None):
        return self._s


maison = {"id": "m1", "nom": "MEMOIRE Complet Vierge.docx", "reference": "synology:/home/Drive/AO/4 - Soumissions/MEMOIRE Complet Vierge.docx"}
cctp = {"id": "c1", "nom": "CCTP LOT 10 - CARRELAGE - FAIENCES.pdf", "reference": "synology:/home/Drive/AO/CCTP.pdf"}
base_db = types.ModuleType("database.connection")


class _PasDeBase:
    async def __aenter__(self):
        raise RuntimeError("pas de base dans le banc")

    async def __aexit__(self, *a):
        return False


base_db.get_db = lambda: _PasDeBase()
sys.modules.setdefault("database", types.ModuleType("database"))
sys.modules["database.connection"] = base_db
espace = {"asyncio": asyncio, "re": re, "logger": logging.getLogger("banc"), "hashlib": __import__("hashlib")}
fonctions("skills/documents_dossier.py",
          {"_imposer_modele", "_mots_forts", "_demande_courante", "_MOTS_CREUX", "_MODELE_MAISON", "_PIECE_CONSULTATION", "_MARQUES_HORS_DEMANDE"}, espace)
espace["dossiers"] = _Dossiers([maison, cctp])
imposer = espace["_imposer_modele"]


def essai(demande, proposee=None):
    return asyncio.run(imposer("u", "f", demande, ["m1", "c1"], proposee))


r = essai("Rôle : chargé d'études. Tâche : Produis une fiche de lecture du DCE : objet du marché, lots, exigences techniques sols.")
verifier("une FICHE DE LECTURE ne prend pas le mémoire vierge pour trame : présentation seule",
         r[0] == "m1" and r[2] is True, str(r))
r = essai("Remplis notre mémoire technique à partir du dossier de la consultation.")
verifier("« remplis notre mémoire » : le mémoire vierge du dossier EST la trame", r[0] == "m1" and r[2] is False, str(r))
r = essai("Rédige la réponse technique du lot 10.", "MEMOIRE Complet Vierge")
verifier("proposé par le modèle de langage : trame", r[0] == "m1" and r[2] is False, str(r))
espace["dossiers"] = _Dossiers([cctp])
r = essai("Produis une fiche de lecture du DCE.")
verifier("sans Word maison dans les pièces : aucun modèle", r[0] is None and r[2] is False, str(r))

verifier("le contrat garde la présentation seule, le plan la suit, la mise en page ne greffe ni garde ni rubriques",
         "data['modele_presentation']=True" in composeur and "plan['modele_presentation']=True" in composeur
         and "if not plan.get('modele_presentation') and (await asyncio.to_thread(_structure,original))['sections']:" in composeur
         and "and not contrat.get('modele_presentation')" in composeur)

# ── 2. Chaque rubrique ne rédige que la sienne, et se lit facilement ──────────
verifier("le rédacteur d'une rubrique n'écrit que SON périmètre (ni la synthèse ni le contrôle d'une autre)",
         "N’écris QUE le contenu de CETTE rubrique" in composeur)
verifier("… et reçoit une consigne de lisibilité (une idée par paragraphe, listes, cellules courtes)",
         "une idée par paragraphe" in composeur and "bloc liste" in composeur)

# ── 3. Les tableaux : colonnes proportionnées ─────────────────────────────────
rendu_src = BACKEND / "bureautique" / "rendu.py"
esp = {}
fonctions("bureautique/rendu.py", {"largeurs_colonnes"}, esp)
largeurs = esp["largeurs_colonnes"]
entetes = ["Rubrique", "Exigence", "Pièce source", "Page"]
lignes = [["Objet du marché", "« Réhabilitation /construction de la piscine de l'ALSH Mouriscot » ; « REHABILITATION DE LA PISCINE DE L'ALSH de MOURISCOT à BIARRITZ (64) » ; description : « démolition d'une piscine extérieure existante et du bâtiment »", "Acte d'engagement LOT 10 ; CCTP LOT 00 - GENERALITES", "Acte d'engagement p.6 ; CCTP LOT 00 p.4"],
          ["Pénalités", "Information absente des pièces fournies.", "—", "—"],
          ["Délais", "« Durée globale des travaux : 12mois » ; « Les travaux s’incorporeront dans le délai global »", "PM.03 - PLANNING TRAVAUX", "PM.03 p.1-2"]]
parts = largeurs(entetes, lignes, 4)
verifier("les parts font toute la largeur", abs(sum(parts) - 1) < 1e-6, str(parts))
verifier("la colonne des citations prend la plus grande part (≥ 35 %)", parts[1] == max(parts) and parts[1] >= 0.35, str([round(p, 2) for p in parts]))
verifier("« Page » reste étroite (≤ 22 %) sans descendre sous 8 %", 0.08 <= parts[3] <= 0.22, str([round(p, 2) for p in parts]))
p2 = largeurs(["A", "B"], [], 2)
verifier("un tableau vide garde des parts sûres", abs(sum(p2) - 1) < 1e-6 and min(p2) > 0, str(p2))

try:
    import docx  # noqa: F401
except ImportError:
    docx = None
    print("  · python-docx absent ici : le rendu Word n'est pas exécuté (il l'est dans l'image du backend)")
if docx is not None:
    spec = importlib.util.spec_from_file_location("bureautique.rendu", rendu_src)
    sys.modules.setdefault("bureautique", types.ModuleType("bureautique"))
    rendu = importlib.util.module_from_spec(spec)
    try:
        spec.loader.exec_module(rendu)
        charge = True
    except Exception as e:  # noqa: BLE001 — module lourd : le banc le dit
        charge = False
        print("  · rendu.py non chargeable ici :", type(e).__name__, str(e)[:120])
    if charge:
        from docx.oxml.ns import qn
        d = docx.Document()
        rendu._tableau_docx(d, {"entetes": entetes, "lignes": lignes}, "1F5FAE")
        t = d.tables[0]
        grille = [c.width for c in t.columns]
        verifier("Word : la grille du tableau suit les parts (colonnes inégales)", len(set(grille)) > 1 and grille[1] == max(grille), str(grille))
        cellule = t.rows[1].cells[1]
        verifier("Word : une cellule qui enchaîne des citations les pose une par ligne, sans rien perdre",
                 len(cellule.paragraphs) == 3 and "".join(p.text for p in cellule.paragraphs).replace(" ;", " ; ").count("«") == lignes[0][1].count("«"),
                 str([p.text[:40] for p in cellule.paragraphs]))
        verifier("Word : les cellules ont des marges intérieures", t._tbl.tblPr.find(qn("w:tblCellMar")) is not None)
        verifier("Word : une cellule courte reste d'un seul tenant", len(t.rows[2].cells[1].paragraphs) == 1)

# ── 4. Les anciens Word (.doc) se lisent ──────────────────────────────────────
verifier("le rédacteur ne met plus les .doc de côté", "'.doc'" in composeur and "DWG, .doc" not in composeur)
import shutil
binaire = shutil.which("soffice") or shutil.which("libreoffice")
if not binaire or docx is None:
    print("  · LibreOffice ou python-docx absent ici : la lecture d'un .doc n'est pas exécutée")
else:
    import subprocess
    import tempfile
    li = importlib.util.spec_from_file_location("bureautique.lecture_integrale", BACKEND / "bureautique" / "lecture_integrale.py")
    lecture = importlib.util.module_from_spec(li); li.loader.exec_module(lecture)
    with tempfile.TemporaryDirectory() as dossier:
        w = docx.Document(); w.add_paragraph("Date limite de remise des offres : le 12 octobre 2026 à 12 h 00.")
        w.save(f"{dossier}/rc.docx")
        subprocess.run([binaire, "--headless", "--convert-to", "doc", "--outdir", dossier, f"{dossier}/rc.docx"], capture_output=True, timeout=120)
        octets = pathlib.Path(f"{dossier}/rc.doc").read_bytes()
    texte = lecture.lire("RC.doc", octets)
    verifier("un vrai .doc (RC) se lit : la date limite ressort", "12 octobre 2026" in texte, texte[:160])

print(f"\n{'═' * 70}\n{'✗ ' + str(len(echecs)) + ' échec(s) : ' + ', '.join(echecs) if echecs else '✓ 0 échec'}\n")
sys.exit(1 if echecs else 0)
