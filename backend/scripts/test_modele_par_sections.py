"""
Banc du MODÈLE REMPLI PAR SECTIONS, de l'en-tête d'un document terminé, de la
file de rédaction et de la réflexion bridée (17/09).

Le relevé : « il a tendance soit à faire un autre document sans reprendre la
trame, soit à reprendre la trame mais à ne modifier que le titre » ; « mets
l'en-tête de la maison » → « présentation inchangée » ; un quantitatif en
« démarrage en attente » pendant que le chat disait « je lance ».

Tout est EXÉCUTÉ sur de vrais fichiers Word fabriqués ici (aucune donnée de
client) : le modèle a une garde, un en-tête avec logo, des rubriques
d'entreprise avec tableau et image, et une rubrique de l'ancien chantier.

    python backend/scripts/test_modele_par_sections.py backend
"""
import asyncio
import io
import os
import struct
import sys
import tempfile
import types
import zipfile
import zlib
from pathlib import Path

BACKEND = Path(sys.argv[1] if len(sys.argv) > 1 else "backend").resolve()
sys.path.insert(0, str(BACKEND))
ECHECS = []


def verifier(nom, condition, detail=""):
    print(("  ✓ " if condition else "  ✗ ") + nom + (f"  → {detail}" if detail and not condition else ""))
    if not condition:
        ECHECS.append(nom)


try:
    from docx import Document
    from docx.shared import Cm
except ImportError:
    print("python-docx absent : banc SAUTÉ (à jouer dans le conteneur ou un venv complet)")
    sys.exit(0)


def _png(couleur=b"\xff\x00\x00"):
    brut = (b"\x00" + couleur * 8) * 8
    def bloc(t, d):
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + bloc(b"IHDR", struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0))
            + bloc(b"IDAT", zlib.compress(brut)) + bloc(b"IEND", b""))


def _modele() -> bytes:
    d = Document()
    d.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Cm(3))
    d.sections[0].footer.paragraphs[0].text = "SAS EXEMPLE-SOLS — capital 12 000 €"
    d.add_paragraph("MEMOIRE TECHNIQUE")
    d.add_paragraph("Projet : Ancien chantier du gymnase")
    d.add_heading("ORGANIGRAMME DE L’ENTREPRISE", level=1)
    d.add_paragraph().add_run().add_picture(io.BytesIO(_png(b"\x00\xff\x00")), width=Cm(2))
    t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = "Année"; t.cell(1, 1).text = "12 salariés"
    d.add_heading("PRINCIPES PROPRES AU CHANTIER", level=1)
    d.add_paragraph("Au gymnase, la pose se fait en trois phases.")
    d.add_paragraph().add_run().add_picture(io.BytesIO(_png(b"\x00\x00\xff")), width=Cm(2))
    d.add_heading("MOYENS TECHNIQUES", level=1)
    d.add_paragraph("Ponceuses, malaxeurs, aspirateurs de classe M.")
    s = io.BytesIO(); d.save(s)
    return s.getvalue()


print("1. La structure du modèle")
from bureautique import sections_modele as sm  # noqa: E402
modele = _modele()
st = sm.structure(modele)
verifier("trois rubriques, dans l'ordre", [s["titre"] for s in st["sections"]] ==
         ["ORGANIGRAMME DE L’ENTREPRISE", "PRINCIPES PROPRES AU CHANTIER", "MOYENS TECHNIQUES"], st["sections"])
verifier("la garde est ce qui précède le premier titre", "MEMOIRE TECHNIQUE" in st["garde"] and "gymnase" in st["garde"])
verifier("images et tableaux comptés par rubrique", st["sections"][0]["images"] == 1 and st["sections"][0]["tableaux"] == 1)
vide = Document(); vide.add_paragraph("sans titre"); tampon = io.BytesIO(); vide.save(tampon)
verifier("un modèle sans titre n'invente pas de découpage", sm.structure(tampon.getvalue())["sections"] == [])

print("2. L'assemblage DANS le modèle")
with tempfile.TemporaryDirectory() as dossier:
    rendu = Document()
    rendu.add_paragraph("Titre fabriqué par le rendu", style="Title")
    titres = ["PRÉSENTATION DE L’OPÉRATION", "ORGANIGRAMME DE L’ENTREPRISE", "MÉTHODE POUR CE PROJET", "MOYENS TECHNIQUES"]
    for titre in titres:
        rendu.add_heading(titre, level=1)
        rendu.add_paragraph("Contenu rédigé pour ce projet : vingt-neuf logements.")
        if titre == "MÉTHODE POUR CE PROJET":     # une illustration dans une rubrique RÉDIGÉE
            rendu.add_picture(io.BytesIO(_png(b"\xff\xff\x00")), width=Cm(2))
    chemin = os.path.join(dossier, "rendu.docx"); rendu.save(chemin)
    bilan = sm.assembler(modele, chemin, titres, {"ORGANIGRAMME DE L’ENTREPRISE": 0, "MOYENS TECHNIQUES": 2})
    final = Document(chemin)
    texte = "\n".join(p.text for p in final.paragraphs)
    verifier("le compte rendu dit 2 reprises et 2 rédigées", bilan["rubriques_reprises"] == 2 and bilan["rubriques_redigees"] == 2, bilan)
    verifier("la garde du modèle est conservée", "MEMOIRE TECHNIQUE" in texte)
    verifier("la garde fabriquée par le rendu est écartée", "fabriqué par le rendu" not in texte)
    verifier("la rubrique d'entreprise reprise garde son tableau", any("12 salariés" in c.text for t in final.tables for r in t.rows for c in r.cells))
    verifier("le texte de remplissage d'une rubrique reprise n'entre pas", texte.count("Contenu rédigé pour ce projet") == 2)
    verifier("la rubrique de l'ANCIEN chantier a disparu", "gymnase, la pose" not in texte)
    ordre = [p.text.strip() for p in final.paragraphs if p.text.strip() in titres]
    verifier("les rubriques suivent l'ordre du PLAN", ordre == titres, ordre)
    z = zipfile.ZipFile(chemin)
    medias = [n for n in z.namelist() if n.startswith("word/media")]
    verifier("l'image de l'ancien chantier ne reste pas cachée dans le fichier", bilan["images_ecartees"] >= 1 and len(medias) == 3, medias)
    verifier("l'archive produite est saine", z.testzip() is None)
    verifier("l'en-tête du modèle (logo) est intact", len(final.sections[0].header._element.findall(
        ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing")) == 1)

print("3. L'en-tête d'un Word posé sur un autre")
cible = Document(); cible.add_paragraph("Mémoire déjà livré."); tampon = io.BytesIO(); cible.save(tampon)
habille, poses = sm.copier_entete(tampon.getvalue(), modele)
resultat = Document(io.BytesIO(habille))
verifier("en-tête et pied recopiés, image comprise", poses["entete"] >= 1 and poses["pied"] >= 1 and poses["images"] == 1, poses)
verifier("le pied porte les mentions de la maison", "EXEMPLE-SOLS" in resultat.sections[0].footer.paragraphs[0].text)
verifier("le corps n'a pas bougé", resultat.paragraphs[0].text == "Mémoire déjà livré.")

print("4. Le skill habiller_document et le refus d'ajouter_document")
with tempfile.TemporaryDirectory() as dossier:
    os.environ["DOCUMENTS_DIR"] = dossier
    for nom in [m for m in sys.modules if m.startswith("bureautique.atelier")]:
        del sys.modules[nom]
    from bureautique import atelier
    jeton = atelier.deposer_fichier("Mémoire.docx", tampon.getvalue(), "u1", "reproduction", "fil-1")
    attaches = types.ModuleType("mail.attaches")
    async def resoudre(refs, user, boite, plafond=None):
        return [{"nom": "entete maison.docx", "octets": modele}], []
    attaches.resoudre = resoudre
    ancien = sys.modules.get("mail.attaches"); sys.modules["mail.attaches"] = attaches
    try:
        from skills import habillage
        user = types.SimpleNamespace(id="u1", email="a@exemple-sols.fr")
        r = asyncio.run(habillage.habiller_document({"entete_de": "entete maison.docx", "_fil": "fil-1"}, user))
        verifier("sans document_id : le dernier Word du fil est habillé", r["version_precedente"] == jeton)
        verifier("une NOUVELLE version est rendue, avec sa carte garantie", r["document_id"] != jeton and r["bloc_ui"]["type"] == "fichier" and r["bloc_garanti"])
        verifier("l'ancienne version reste disponible", bool(atelier.chemin_fichier(jeton, "u1")))
    finally:
        if ancien is not None:
            sys.modules["mail.attaches"] = ancien
        else:
            sys.modules.pop("mail.attaches", None)
bureau = (BACKEND / "skills" / "bureau.py").read_text(encoding="utf-8")
verifier("ajouter_document REFUSE la présentation d'un document terminé, et nomme le bon geste",
         'actuelle.get("fini"):' in bureau and "habiller_document" in bureau)
familles = (BACKEND / "skills" / "familles.py").read_text(encoding="utf-8")
verifier("le nouveau geste est dans la famille des documents", '"habiller_document"' in familles)

print("5. Le choix du modèle et le plan")
src = (BACKEND / "skills" / "documents_dossier.py").read_text(encoding="utf-8")
import ast, re, json, hashlib, logging  # noqa: E401,E402
arbre = ast.parse(src)
noms = {"_mots_forts", "_normaliser_plan"}
corps = [n for n in arbre.body if (isinstance(n, ast.FunctionDef) and n.name in noms)
         or (isinstance(n, ast.Assign) and getattr(n.targets[0], "id", "") in ("_PIECE_CONSULTATION", "_MODELE_MAISON", "_MOTS_CREUX"))]
esp = {"re": re, "json": json, "hashlib": hashlib, "logger": logging.getLogger("banc")}
exec(compile(ast.Module(body=corps, type_ignores=[]), "dossier_extrait", "exec"), esp)
verifier("le cadre de réponse du maître d'ouvrage est une PIÈCE DE CONSULTATION",
         bool(esp["_PIECE_CONSULTATION"].search("Cadre de réponse technique.docx")) and not esp["_PIECE_CONSULTATION"].search("MEMOIRE Complet Vierge.docx"))
verifier("le mémoire vierge de l'entreprise est un modèle maison", bool(esp["_MODELE_MAISON"].search("MEMOIRE Complet Vierge.docx")))
verifier("« mémoire technique type » partage un mot fort avec « fais-moi un mémoire technique »",
         bool(esp["_mots_forts"]("mémoire technique type") & esp["_mots_forts"]("peux-tu me faire un mémoire technique")))
plan = {"sections": [{"titre": "Notre organigramme", "sources": [], "reprise_modele": 0},
                     {"titre": "Méthode", "sources": ["a"], "reprise_modele": 9},
                     {"titre": "Doublon", "sources": [], "reprise_modele": 0}], "sources_ecartees": {}}
norme = esp["_normaliser_plan"](plan, ["a", "m"], "m", st)
verifier("une rubrique reprise prend le TITRE EXACT du modèle", norme["sections"][0]["titre"] == "ORGANIGRAMME DE L’ENTREPRISE")
verifier("un index inconnu ou déjà repris redevient une rubrique rédigée",
         "reprise_modele" not in norme["sections"][1] and "reprise_modele" not in norme["sections"][2])
verifier("le modèle imposé n'a pas à être « affecté à une rubrique »", norme["modele_source"] == "m" and "m" in norme["sources_ecartees"])
verifier("une rubrique reprise ne se RÉDIGE pas (aucun appel de modèle)", "if type(section.get('reprise_modele')) is int:" in src)
verifier("le corps du modèle d'origine accompagne le rendu", "_modele_original" in src and "_modele_original" in (BACKEND / "bureautique" / "rendu.py").read_text(encoding="utf-8"))
trames = (BACKEND / "skills" / "trames.py").read_text(encoding="utf-8")
verifier("utiliser_trame sans remplacements rend les rubriques et nomme le geste qui REMPLIT", "composer_document_dossier" in trames and '"rubriques"' in trames)
illus = (BACKEND / "bureautique" / "illustrations.py").read_text(encoding="utf-8")
verifier("une illustration refusée est sautée, elle n'arrête plus la rédaction", "Illustration sautée" in illus)

print("6. La file de rédaction : pas de doublon, et la position se dit")
with tempfile.TemporaryDirectory() as dossier:
    os.environ["DOCUMENTS_DIR"] = dossier
    for nom in [m for m in sys.modules if m.startswith(("ressources.", "bureautique.atelier"))]:
        del sys.modules[nom]
    from ressources import documents_file as fd, dossiers
    dossiers.enregistrer("u1", "fil-A", "cctp.pdf", "texte du cctp " * 30, "ref", None)
    dossiers.enregistrer("u1", "fil-B", "cctp.pdf", "texte du cctp " * 30, "ref", None)
    a = fd.soumettre("u1", "fil-A", "quantitatif", {"demande": "Métrés du projet", "_demande_utilisateur": "fais les métrés du projet"})
    b = fd.soumettre("u1", "fil-B", "quantitatif", {"demande": "Métrés du projet (bis)", "_demande_utilisateur": "Fais  les métrés du projet"})
    verifier("la même demande, tapée dans une autre conversation, ne crée pas un second travail",
             b.get("deja_lance") is True and b["tache_documentaire"] == a["tache_documentaire"], b)
    c = fd.soumettre("u1", "fil-B", "document", {"demande": "Rédige le mémoire", "_demande_utilisateur": "rédige le mémoire"})
    verifier("un travail mis en file DIT combien passent avant lui", c.get("travaux_devant") == 1 and "en file" in c["a_faire"].casefold(), c)
    phases = [p["phase"] for p in fd.progression("u1", "fil-B")]
    verifier("l'avancement dit la position, plus « démarrage en attente »",
             any("En file : 1 travail" in p for p in phases) and not any("démarrage en attente" in p for p in phases), phases)

print("7. La réflexion bridée, d'après la mesure")
routeur_src = (BACKEND / "llm" / "router.py").read_text(encoding="utf-8")
arbre = ast.parse(routeur_src)
corps = [n for n in arbre.body if (isinstance(n, ast.FunctionDef) and n.name in ("reflexion_mesuree", "texte_seul"))
         or (isinstance(n, ast.Assign) and getattr(n.targets[0], "id", "") in ("_TEXTE_SEUL",))]
from typing import Optional  # noqa: E402
esp = {"Optional": Optional, "settings": types.SimpleNamespace(ollama_cloud_reflexion="mesuree")}
exec(compile(ast.Module(body=corps, type_ignores=[]), "routeur_extrait", "exec"), esp)
r = esp["reflexion_mesuree"]
verifier("deepseek flash : jamais de réflexion (mesuré : 5,6 s au lieu d'une réponse vide)",
         r("ollama_cloud", "deepseek-v4-flash:0731", "standard") == "none" and r("ollama_cloud", "deepseek-v4.1-flash", "complex") == "none")
verifier("deepseek pro : réflexion basse au palier complexe seulement",
         r("ollama_cloud", "deepseek-v4-pro:0813", "complex") == "low" and r("ollama_cloud", "deepseek-v4-pro:0813", "standard") == "none")
verifier("kimi et glm : « low » (« none » les fait penser en clair)", r("ollama_cloud", "kimi-k3") == "low" and r("ollama_cloud", "glm-5.3-flash") == "low")
verifier("la vision de deepseek-v4.1-flash se bride en « low »", r("ollama_cloud", "deepseek-v4.1-flash", usage="vision") == "low")
verifier("rien n'est envoyé aux autres fournisseurs ni aux modèles non mesurés",
         r("openrouter", "deepseek-v4-flash") is None and r("ollama_cloud", "gpt-oss:120b") is None)
esp["settings"].ollama_cloud_reflexion = "libre"
verifier("OLLAMA_CLOUD_REFLEXION=libre rend la main aux modèles", r("ollama_cloud", "deepseek-v4-flash:0731") is None)
verifier("deepseek-v4.1-flash VOIT : il n'est plus écarté sur son nom",
         not esp["texte_seul"]("deepseek-v4.1-flash") and esp["texte_seul"]("deepseek-v4-flash:0731") and esp["texte_seul"]("deepseek-v4-pro:0813"))

reglages_src = (BACKEND / "llm" / "reglages.py").read_text(encoding="utf-8")
verifier("le réglage « Vision » accepte lui aussi deepseek-v4.1-flash (même liste que la cascade)",
         '"deepseek-v4-flash", "deepseek-v4-pro"' in reglages_src and '("deepseek-v4", "deepseek-chat"' not in reglages_src)

print("8. L'écran et l'exploitation")
chat = (BACKEND.parent / "frontend" / "components" / "chat" / "ChatWindow.tsx").read_text(encoding="utf-8")
verifier("l'avancement des rédactions vit SOUS la réflexion, plus en haut du fil",
         chat.index("<ReflexionEnCours") < chat.index("<SuiviRedactions") < chat.index("<InputBar")
         and chat.index("<MessageList") < chat.index("<SuiviRedactions"))
verifier("le bandeau « Travail en cours » du haut a disparu", "<SuiviTravail" not in chat)
requetes = (BACKEND / "agents" / "requetes.py").read_text(encoding="utf-8")
verifier("le signe de vie ne peut plus ANNULER le tour qu'il surveille", "parent.cancel()" not in requetes)
nginx = (BACKEND.parent / "nginx" / "nginx.conf").read_text(encoding="utf-8")
verifier("nginx transmet l'adresse posée par Caddy, pas celle de Caddy",
         "X-Forwarded-For $remote_addr;" not in nginx and nginx.count("X-Forwarded-For $proxy_add_x_forwarded_for;") >= 8)
ignore = (BACKEND / ".dockerignore").read_text(encoding="utf-8")
verifier("les modes d'emploi entrent dans l'image", "!outils/docs/*.md" in ignore)

print("9. Le quantitatif : pièces rejointes, contexte borné, suspension qui tient")
# Relevé du 17/09 : « Source inconnue ou ambiguë » — le DPGF du lot 11 figurait
# CINQ fois dans le dossier (un exemplaire par envoi du même fichier) ; et chaque
# appel du quantitatif portait 181 666 caractères, donc expirait.
with tempfile.TemporaryDirectory() as dossier:
    os.environ["DOCUMENTS_DIR"] = dossier
    for nom in [m for m in sys.modules if m.startswith(("ressources.", "bureautique.atelier"))]:
        del sys.modules[nom]
    from ressources import documents_file as fd, dossiers
    contenu = "Feuille DPGF — ligne 4 : A4=\"Carrelage 45x45\" | B4=\"1250\" | C4=\"m²\"\n" * 20
    a = dossiers.enregistrer("u1", "fil-Q", "DPGF-LOT 11 CARRELAGE.xlsx", contenu, "/api/documents/jeton-1", "sha-dpgf")
    b = dossiers.enregistrer("u1", "fil-Q", "DPGF-LOT 11 CARRELAGE.xlsx", contenu, "/api/documents/jeton-2", "sha-dpgf")
    verifier("le même fichier rejoint reste UNE pièce, sa référence se rafraîchit",
             a == b and dossiers.sources("u1", "fil-Q", [a])[0]["reference"].endswith("jeton-2"))
    verifier("le manifeste ne la compte qu'une fois", len(dossiers.manifeste("u1", "fil-Q")) == 1)
    verifier("un nom raccourci par le modèle se résout (« DPGF-LOT 11 »)",
             dossiers.sources("u1", "fil-Q", ["DPGF-LOT 11"])[0]["id"] == a)
    dossiers.enregistrer("u1", "fil-Q", "Devis.pdf", "devis du client A, 1 200 m² de sol " * 5, "/api/documents/x", "sha-A")
    dossiers.enregistrer("u1", "fil-Q", "Devis.pdf", "devis du client B, 300 m² de faïence " * 5, "/api/documents/y", "sha-B")
    try:
        dossiers.sources("u1", "fil-Q", ["Devis.pdf"]); refus = ""
    except ValueError as e:
        refus = str(e)
    verifier("deux fichiers DIFFÉRENTS du même nom restent une ambiguïté, et le refus donne les identifiants",
             "Plusieurs pièces" in refus and "IDENTIFIANT" in refus, refus)
    verifier("…et ils restent deux pièces au manifeste", len(dossiers.manifeste("u1", "fil-Q")) == 3)
    # la suspension tient, « retirer » masque sans effacer
    r = fd.soumettre("u1", "fil-Q", "quantitatif", {"demande": "métrés", "_demande_utilisateur": "fais les métrés"})
    cle = r["tache_documentaire"]
    fd.piloter("u1", "fil-Q", cle, reprendre=False)
    fd._maj(cle, statut="attente", prochain=0)          # ce que fait l'essai en cours en se terminant
    verifier("une suspension n'est PAS écrasée par l'essai qui se termine",
             [p["statut"] for p in fd.progression("u1", "fil-Q")] == ["suspendu"])
    fd.piloter("u1", "fil-Q", cle, retirer=True)
    verifier("« retirer » masque le travail du suivi", fd.progression("u1", "fil-Q") == [])
    with dossiers.base() as c:
        verifier("…sans rien effacer en base", c.execute("select statut from file_documentaire where id=?", (cle,)).fetchone()[0] == "retire")
quanti = (BACKEND / "skills" / "quantitatifs.py").read_text(encoding="utf-8")
arbre = ast.parse(quanti)
corps = [n for n in arbre.body if isinstance(n, ast.Assign) and getattr(n.targets[0], "id", "") in ("_QUANTITE", "PLAFOND_CONTEXTE")]
esp = {"re": re}
exec(compile(ast.Module(body=corps, type_ignores=[]), "quanti_extrait", "exec"), esp)
verifier("un passage de métré porte une quantité (« 1 250 m² », « 42,5 ml », « 12 u »)",
         all(esp["_QUANTITE"].search(x) for x in ("Carrelage 45x45 : 1 250 m²", "plinthes 42,5 ml", "siphons : 12 u")))
verifier("une clause administrative n'en porte pas : elle n'est pas interrogée",
         not esp["_QUANTITE"].search("Le titulaire remet son mémoire avant la date limite fixée au règlement."))
verifier("le contexte commun est BORNÉ et ne transmet plus les analyses entières",
         esp["PLAFOND_CONTEXTE"] <= 20000 and "'affectations':contexte}" not in quanti and "contexte_pour(s)" in quanti)
suivi = (BACKEND.parent / "frontend" / "components" / "chat" / "SuiviRedactions.tsx").read_text(encoding="utf-8")
verifier("l'écran propose « retirer » sur un travail arrêté", '"retirer"' in suivi and "Retirer ce travail du suivi" in suivi)

print(("✗ %d échec(s) : %s" % (len(ECHECS), ", ".join(ECHECS))) if ECHECS else "✓ 0 échec")
sys.exit(1 if ECHECS else 0)
