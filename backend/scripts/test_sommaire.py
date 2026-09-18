"""
LE SOMMAIRE D'UN MODÈLE REMPLI (18/09, banc Duret, Q19).

« Refais le mémoire uniquement pour le lot 12 » : corps juste (sols souples seuls),
mais le sommaire de la trame listait encore « 1) Carrelages et faïence »,
« f) Pose du carrelage au sol » — son cache figé. `bureautique/sommaire.py` le
refait d'après les titres réels et leur page rendue.

Le banc fabrique un Word de la forme de la trame (champ TOC ouvert dans la
première ligne, lignes TM1/TM2 à liens et PAGEREF imbriqués, fermeture dans son
propre paragraphe), et un rendu PDF DOUBLÉ par PyMuPDF (LibreOffice n'est pas sur
ce poste) : pages de garde, de sommaire, puis corps avec les numéros affichés.
"""
import pathlib
import sys
import tempfile

BACKEND = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "backend").resolve()
sys.path.insert(0, str(BACKEND))
echecs = []


def verifier(nom, cond, detail=""):
    print(("  ✓ " if cond else "  ✗ ") + nom + ("" if cond else f"  → {detail}"))
    if not cond:
        echecs.append(nom)


try:
    import fitz  # noqa: F401
    from docx import Document
    from docx.oxml import parse_xml
except ImportError as e:
    print(f"SKIP : {e}")
    sys.exit(0)

from bureautique import sommaire

print("═══ SOMMAIRE D'UN MODÈLE REMPLI — " + str(BACKEND))
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def ligne(style, label, titre, page, ancre, ouvrir=False):
    debut = ('<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
             '<w:r><w:fldChar w:fldCharType="separate"/></w:r>') if ouvrir else ""
    return parse_xml(
        f'<w:p {W}><w:pPr><w:pStyle w:val="{style}"/><w:tabs><w:tab w:val="left" w:pos="440"/><w:tab w:val="right" w:leader="dot" w:pos="9062"/></w:tabs></w:pPr>'
        f'{debut}<w:hyperlink w:anchor="{ancre}"><w:r><w:rPr><w:noProof/></w:rPr><w:t>{label}</w:t></w:r><w:r><w:tab/></w:r><w:r><w:rPr><w:noProof/></w:rPr><w:t>{titre}</w:t></w:r><w:r><w:tab/></w:r>'
        f'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGEREF {ancre} \\h </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        f'<w:r><w:t>{page}</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:hyperlink></w:p>')


doc = Document()
corps = doc.element.body
doc.add_paragraph("MÉMOIRE TECHNIQUE — garde")
p_titre = doc.add_paragraph("Table des matières")
for i, (st, lab, t, pg) in enumerate([("TM1", "I.", "ORGANIGRAMME", 3), ("TM1", "II.", "Principes de réalisations", 3),
                                      ("TM2", "1)", "Carrelages et faïence", 3), ("TM2", "2)", "Sols souples", 4)]):
    corps.insert(corps.index(p_titre._p) + 1 + i, ligne(st, lab, t, pg, f"_Toc{i}", ouvrir=(i == 0)))
fin = parse_xml(f'<w:p {W}><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')
corps.insert(corps.index(p_titre._p) + 5, fin)
doc.add_heading("ORGANIGRAMME", level=1)
doc.add_paragraph("L'organigramme.")
doc.add_heading("Principes de réalisations", level=1)
doc.add_heading("Nettoyage et préparation", level=2)
doc.add_paragraph("Texte.")
doc.add_heading("Collage du revêtement", level=2)
doc.add_heading("Points à confirmer", level=1)

dossier = tempfile.mkdtemp()
chemin = f"{dossier}/memoire.docx"
doc.save(chemin)


def faux_pdf(pages):
    d = fitz.open()
    for texte in pages:
        pg = d.new_page()
        pg.insert_text((50, 60), texte, fontsize=10)
    return d.tobytes()


# Le rendu : garde, sommaire (lignes à points de conduite), corps numéroté par le traitement de texte.
RENDU = ["MÉMOIRE TECHNIQUE — garde",
         "Table des matières\nI.\nORGANIGRAMME........................3\nII.\nPrincipes de réalisations........3\n1)\nCarrelages et faïence............3\n2)\nSols souples.....................4",
         "I.\nORGANIGRAMME\nL'organigramme.",
         "II.\nPrincipes de réalisations\n1) Nettoyage et préparation\nTexte.",
         "2) Collage du revêtement\nIII.\nPoints à confirmer"]
appels = []


def convertir(c):
    appels.append(c)
    return faux_pdf(RENDU)


r = sommaire.actualiser(chemin, convertir=convertir)
verifier("le sommaire est actualisé, chaque titre avec sa page", r.get("sommaire") == "actualisé" and r.get("sans_page") == 0, r)
d2 = Document(chemin)
from docx.oxml.ns import qn
lignes = [(p.find("./" + qn("w:pPr") + "/" + qn("w:pStyle")).get(qn("w:val")) if p.find("./" + qn("w:pPr") + "/" + qn("w:pStyle")) is not None else None,
           "|".join(t.text for t in p.iter(qn("w:t"))))
          for p in d2.element.body.iter(qn("w:p"))]
texte = " ".join(t for _, t in lignes)
verifier("les lignes de l'ANCIEN chantier ont quitté le sommaire", "Carrelages et faïence" not in texte and "Sols souples" not in texte, texte[:300])
entrees = [(s, t) for s, t in lignes if s in ("TM1", "TM2")]
verifier("une ligne par titre réel, au style de son niveau (TM1 / TM2)",
         [s for s, _ in entrees] == ["TM1", "TM1", "TM2", "TM2", "TM1"], entrees)
verifier("numéro, titre et page lus au rendu (« II.|Principes de réalisations|4 », « 2)|Collage du revêtement|5 »)",
         entrees[1][1] == "II.|Principes de réalisations|4" and entrees[3][1] == "2)|Collage du revêtement|5", entrees)
debut, _ = sommaire._champ_du_sommaire(d2.element.body, qn)
types = [fc.get(qn("w:fldCharType")) for fc in d2.element.body.iter(qn("w:fldChar"))]
verifier("le champ reste un champ, équilibré : ouvert dans la 1re ligne, fermé dans son paragraphe",
         debut is not None and types.count("begin") == types.count("end") == 1 and sommaire._fin_du_champ(debut, qn) is not None, types)
verifier("le rendu a été relu une seconde fois seulement si les pages bougeaient (ici : 2 rendus au plus)", 1 <= len(appels) <= 2, len(appels))

# Sans LibreOffice : le fichier reste tel quel, jamais perdu.
avant = open(chemin, "rb").read()


def en_panne(c):
    raise ValueError("Le moteur LibreOffice n’est pas installé")


r2 = sommaire.actualiser(chemin, convertir=en_panne)
verifier("sans moteur de rendu : « non actualisé », fichier intact", r2.get("sommaire") == "non actualisé" and open(chemin, "rb").read() == avant, r2)

# Un Word sans sommaire n'est pas touché.
nu = Document(); nu.add_heading("Titre", level=1); nu.save(f"{dossier}/nu.docx")
verifier("un document sans sommaire n'est pas touché", sommaire.actualiser(f"{dossier}/nu.docx", convertir=convertir) == {"sommaire": "absent"})

# ── La numérotation des sous-titres repart dans chaque rubrique (même geste que la trame) ──
print("\nNumérotation des sous-titres rédigés")
from bureautique.sections_modele import _relancer_la_numerotation
n = Document()
num = n.part.numbering_part.element
num.append(parse_xml(f'<w:abstractNum {W} w:abstractNumId="90"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="upperRoman"/></w:lvl></w:abstractNum>'))
num.append(parse_xml(f'<w:abstractNum {W} w:abstractNumId="91"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/></w:lvl></w:abstractNum>'))
premier_num = num.find(qn("w:num"))
for nid, aid in (("46", "90"), ("48", "91"), ("49", "91")):
    el = parse_xml(f'<w:num {W} w:numId="{nid}"><w:abstractNumId w:val="{aid}"/></w:num>')
    num.append(el)
for style, nid in (("Heading1", "46"), ("Heading3", "48")):
    st = next(x for x in n.styles.element.iter(qn("w:style")) if x.get(qn("w:styleId")) == style)
    ppr = st.find(qn("w:pPr"))
    if ppr is None:
        ppr = parse_xml(f"<w:pPr {W}/>"); st.append(ppr)
    ppr.append(parse_xml(f'<w:numPr {W}><w:numId w:val="{nid}"/></w:numPr>'))
n.add_heading("CHANTIERS", level=1)
a1 = n.add_heading("Réunion de lancement", level=3)
a1._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {W}><w:ilvl w:val="0"/><w:numId w:val="49"/></w:numPr>'))
n.add_heading("Mises au point", level=3)
n.add_heading("PERSONNEL ET OUVRIERS", level=1)
b1 = n.add_heading("Personnel", level=3)
b2 = n.add_heading("Planning", level=3)
relancees = _relancer_la_numerotation(n)
np_b1 = b1._p.find("./" + qn("w:pPr") + "/" + qn("w:numPr"))
verifier("une seule relance : le premier sous-titre de la rubrique rédigée (celle de la trame est déjà là)",
         relancees == 1 and np_b1 is not None and b2._p.find("./" + qn("w:pPr") + "/" + qn("w:numPr")) is None, relancees)
nouveau = next(x for x in num.findall(qn("w:num")) if x.get(qn("w:numId")) == np_b1.find(qn("w:numId")).get(qn("w:val")))
verifier("la relance reprend la définition du style (même liste) et remet le départ à 1",
         nouveau.find(qn("w:abstractNumId")).get(qn("w:val")) == "91"
         and nouveau.find("./" + qn("w:lvlOverride") + "/" + qn("w:startOverride")).get(qn("w:val")) == "1")
n.save(f"{dossier}/num.docx"); Document(f"{dossier}/num.docx")
verifier("le document se rouvre", True)

src = (BACKEND / "bureautique" / "rendu.py").read_text(encoding="utf-8")
verifier("le rendu dans la trame actualise le sommaire après l'assemblage", 'bilan["sommaire"] = _sommaire_a_jour(sortie)' in src)

print("\n" + ("✓ 0 échec" if not echecs else f"✗ {len(echecs)} échec(s)"))
sys.exit(1 if echecs else 0)
