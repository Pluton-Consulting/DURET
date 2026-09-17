"""
HABILLER UN DOCUMENT DÉJÀ TERMINÉ — son en-tête et son pied (17/09).

Relevé de Noa (Duret) : « mets en en-tête l'en-tête de Revêtements Duret Sols »
sur un mémoire déjà livré → deux `ajouter_document` avec `entete_image`, deux
fois « ajoutes=0, présentation inchangée », boucle fermée. Trois causes :

  * la présentation ne se règle QUE sur un document ouvert ; sur un document
    terminé le geste ne faisait rien, et ne le disait pas comme un refus ;
  * un rendu terminé est immuable (et c'est voulu) : la seule voie est une
    NOUVELLE version — qu'aucun geste n'ouvrait ;
  * l'en-tête de la maison est un DOCUMENT WORD (logo ancré en haut, mentions
    légales en pied), pas une image : `entete_image` ne pouvait pas le poser.

`habiller_document` prend un Word terminé (ou déposé), recopie l'en-tête et le
pied d'un autre Word — ou pose une image en en-tête —, et rend une NOUVELLE
version téléchargeable. L'original ne bouge pas.
"""
from __future__ import annotations

import asyncio
import io
import logging

from skills.erreurs import SkillError

logger = logging.getLogger("pluton.skills.habillage")

EXT_IMAGES = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")


def _proprietaire(user) -> str:
    return str(getattr(user, "id", "") or "")


def _image_en_entete(document: bytes, image: bytes) -> bytes:
    """Une IMAGE posée en en-tête de chaque section, à la largeur utile."""
    from docx import Document
    from docx.shared import Emu
    doc = Document(io.BytesIO(document))
    for section in doc.sections:
        entete = section.header
        entete.is_linked_to_previous = False
        for enfant in list(entete._element):
            entete._element.remove(enfant)
        utile = section.page_width - section.left_margin - section.right_margin
        entete.add_paragraph().add_run().add_picture(io.BytesIO(image), width=Emu(int(utile)))
    sortie = io.BytesIO()
    doc.save(sortie)
    Document(io.BytesIO(sortie.getvalue()))
    return sortie.getvalue()


async def habiller_document(parametres: dict, utilisateur) -> dict:
    from bureautique import atelier
    from mail.attaches import resoudre

    proprietaire = _proprietaire(utilisateur)
    fil = parametres.get("_fil")
    jeton = str(parametres.get("document_id") or parametres.get("document") or "").strip()
    if not jeton:
        dernier = await asyncio.to_thread(atelier.dernier_livrable_du_fil, proprietaire, fil)
        jeton = str((dernier or {}).get("jeton") or (dernier or {}).get("document_id") or "")
    chemin = await asyncio.to_thread(atelier.chemin_fichier, jeton, proprietaire) if jeton else None
    if not chemin:
        raise SkillError("Aucun document terminé à habiller : donne le `document_id` EXACT d'un Word "
                         "déjà produit dans cette conversation (celui de sa carte).")
    if not chemin.lower().endswith(".docx"):
        raise SkillError("Seul un Word (.docx) reçoit un en-tête ici. Pour un PDF, habille d'abord "
                         "le Word, puis redemande le PDF.")
    designation = (parametres.get("entete_de") or parametres.get("entete") or parametres.get("source")
                   or parametres.get("fichier"))
    if not designation:
        raise SkillError("Dis QUEL en-tête poser : `entete_de` = le nom ou le chemin du Word (ou de "
                         "l'image) qui porte l'en-tête de la maison, tel qu'un geste te l'a rendu.")
    pretes, refusees = await resoudre([designation], utilisateur,
                                      str(getattr(utilisateur, "email", "") or ""),
                                      plafond=40 * 1024 * 1024)
    if not pretes:
        raison = (refusees[0].get("raison") if refusees else "") or "fichier introuvable"
        raise SkillError(f"L'en-tête « {designation} » n'a pas pu être ouvert : {raison}")
    source = pretes[0]
    with open(chemin, "rb") as f:
        document = f.read()

    nom_source = str(source.get("nom") or "")
    if nom_source.lower().endswith(".docx"):
        from bureautique.sections_modele import copier_entete
        try:
            produit, poses = await asyncio.to_thread(copier_entete, document, source["octets"])
        except Exception as e:  # noqa: BLE001 — un Word abîmé se dit, il ne plante pas le tour
            logger.warning("Habillage impossible depuis %s : %s", nom_source, e)
            raise SkillError(f"L'en-tête de « {nom_source} » n'a pas pu être recopié "
                             f"({type(e).__name__}). Le document d'origine est intact.") from e
        if not poses["entete"] and not poses["pied"]:
            raise SkillError(f"« {nom_source} » ne porte ni en-tête ni pied de page : rien à recopier.")
        detail = (f"en-tête et pied de « {nom_source} » recopiés"
                  + (f", {poses['images']} image(s) comprise(s)" if poses["images"] else ""))
    elif nom_source.lower().endswith(EXT_IMAGES):
        produit = await asyncio.to_thread(_image_en_entete, document, source["octets"])
        detail = f"image « {nom_source} » posée en en-tête de chaque page"
    else:
        raise SkillError(f"« {nom_source} » n'est ni un Word ni une image : je ne sais pas en tirer un en-tête.")

    fiche = await asyncio.to_thread(atelier.fiche, jeton, proprietaire) or {}
    titre = str((fiche.get("entete") or {}).get("titre") or fiche.get("nom") or "Document").strip()
    nom_sortie = (titre[:-5] if titre.lower().endswith(".docx") else titre) + ".docx"
    nouveau = await asyncio.to_thread(atelier.deposer_fichier, nom_sortie, produit, proprietaire,
                                      "reproduction", fil)
    return {
        "document_id": nouveau, "version_precedente": jeton, "habillage": detail, "octets": len(produit),
        "bloc_garanti": True,
        "bloc_ui": {"type": "fichier", "nom": nom_sortie, "url": f"/api/documents/{nouveau}",
                    "titre": titre, "format": "docx", "octets": len(produit)},
        "message_final": f"Nouvelle version du document : {detail}. Le contenu n'a pas changé.",
        "a_faire": ("La carte de la nouvelle version s'affiche AUTOMATIQUEMENT : n'écris aucun bloc. "
                    "Dis en une phrase ce qui a été posé. L'ancienne version reste disponible."),
    }


from skills.registre import Declaration  # noqa: E402

SKILLS = {
    "habiller_document": Declaration(
        fonction=habiller_document,
        description=(
            "POSE L'EN-TETE ET LE PIED DE PAGE de la maison sur un Word DEJA TERMINE, et rend une "
            "nouvelle version (l'ancienne reste). `entete_de` : le nom ou le chemin EXACT du Word qui "
            "porte l'en-tete (ex. « entete duret.docx », rendu par une recherche ou un listage), ou "
            "une image. `document_id` : le Word a habiller (sans lui : le dernier document de la "
            "conversation). C'est LE geste pour « mets notre en-tete sur ce document » : "
            "`ajouter_document` ne regle la presentation que d'un document encore OUVERT."),
        requis=["entete_de"],
        optionnels=["document_id"],
        effet="ecriture_interne",
        libelle="je pose l'en-tête de la maison sur le document",
    ),
}
