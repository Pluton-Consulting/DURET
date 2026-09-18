"""Quantitatif traçable : lignes métier reliées aux preuves et calcul décimal.

Aucune quantité issue des pixels n'est promue en mesure exacte. Les postes sans
cotes ou sans correspondance CCTP sont rendus dans les réserves, pas inventés.
"""
import asyncio,hashlib,json,re
from decimal import Decimal
from ressources import dossiers
from skills.registre import Declaration
from skills.documents_dossier import _json,_analyses
from skills.chiffres_sources import calculer,NOMBRE,VALEUR


def nombre(texte):
    return Decimal(str(texte).replace(' ','').replace('\u00a0','').replace('\u202f','').replace(',','.'))

_UNITES_ECRITES={'m2':'m²','m²':'m²','m3':'m³','m³':'m³','ml':'ml','m.l':'ml','ml.':'ml','m':'m','u':'u','un':'u','unité':'u','unite':'u','unités':'u','pce':'u','kg':'kg'}
def _unite(brut):
    return _UNITES_ECRITES.get(str(brut or '').strip().casefold().replace(' ',''),str(brut or '').strip())
def _plat(texte):
    """Pour retrouver une citation : espaces, guillemets de cellule et « m2 »/« m² » ne comptent pas."""
    t=str(texte or '').casefold().replace('m2','m²').replace('m3','m³').replace('"',' ').replace('|',' ')
    return ' '.join(t.split())

def _verifier_ligne(ligne,preuves):
    # SOUPLE SUR LE CADRE, STRICT SUR LE CHIFFRE (17/09). Une ligne sans niveau ou
    # sans local n'est pas fausse : un total de poste n'en a pas. La rejeter
    # faisait perdre des quantités VRAIES ; seuls le poste, la valeur citée,
    # l'unité et le calcul restent non négociables.
    for k in ('lot','poste'):
        if not isinstance(ligne.get(k),str) or not ligne[k].strip():raise ValueError('Ligne quantitative sans '+k)
    for k,defaut in (('niveau','Non précisé'),('local','Non précisé')):
        if not isinstance(ligne.get(k),str) or not ligne[k].strip():ligne[k]=defaut
    ligne['unite']=_unite(ligne.get('unite'))
    for o in (ligne.get('operandes') or []):
        if isinstance(o,dict):o['unite']=_unite(o.get('unite'))
    if ligne['unite'] not in ('m²','m','ml','m³','u','kg'):raise ValueError('Unité de résultat inconnue.')
    valeurs={};references=[]
    operandes=ligne.get('operandes')
    if not isinstance(operandes,list) or not 1<=len(operandes)<=12:raise ValueError('De 1 à 12 opérandes sourcés sont nécessaires.')
    for i,o in enumerate(operandes):
        extrait=preuves.get(o.get('preuve'))
        citation=str(o.get('citation') or '')
        if not extrait or not citation or _plat(citation) not in _plat(extrait):raise ValueError('Citation de quantité absente de la source.')
        valeur=nombre(o.get('valeur'))
        trouves=[]
        for m in NOMBRE.finditer(citation):
            v=VALEUR.match(m.group());trouves.append(nombre(v[1]))
        if valeur not in trouves:raise ValueError('La valeur ne figure pas dans la citation donnée.')
        unite=o.get('unite')
        if unite not in ('m²','m','ml','m³','u','kg'):raise ValueError('Unité d’opérande inconnue ; ne déduis pas une unité absente.')
        # L'unité est citée dans le même extrait (en-tête de tableau compris).
        alias={'m²':('m²','m2'),'m³':('m³','m3'),'ml':('ml','m'),'m':('m',),'u':('unité','unite',' u','quantité','quantite'),'kg':('kg',)}[unite]
        if not any(re.search(r'(?<![\w])'+re.escape(a.strip())+r'(?![\w])',citation.lower()) for a in alias):raise ValueError('L’unité doit apparaître dans la citation, avec la valeur.')
        valeurs['c'+str(i+1)]=str(valeur);references.append(o)
    resultat=Decimal(calculer(ligne.get('formule') or 'c1',valeurs))
    if resultat<0:raise ValueError('Une quantité négative nécessite de revoir les déductions.')
    justification=ligne.get('affectation') or {}
    if not justification.get('preuve') in preuves or not justification.get('citation'):raise ValueError('L’affectation au poste nécessite une preuve du CCTP, DPGF ou tableau de localisation.')
    if _plat(justification['citation']) not in _plat(preuves[justification['preuve']]):raise ValueError('Affectation non prouvée par les pièces.')
    # Vérification dimensionnelle : surfaces additionnées ; longueurs multipliées.
    import ast
    dimensions={'m':(1,0),'ml':(1,0),'m²':(2,0),'m³':(3,0),'u':(0,0),'kg':(0,1)}
    dims={'c'+str(i+1):dimensions[o['unite']] for i,o in enumerate(operandes)}
    def dimension(n):
        if isinstance(n,ast.Name):return dims[n.id]
        if isinstance(n,ast.UnaryOp):return dimension(n.operand)
        if isinstance(n,ast.BinOp):
            a,b=dimension(n.left),dimension(n.right)
            if isinstance(n.op,(ast.Add,ast.Sub)):
                if a!=b:raise ValueError('Addition de grandeurs incompatibles.')
                return a
            if isinstance(n.op,ast.Mult):return tuple(x+y for x,y in zip(a,b))
            if isinstance(n.op,ast.Div):return tuple(x-y for x,y in zip(a,b))
        raise ValueError('Formule dimensionnelle invalide.')
    if dimension(ast.parse(ligne.get('formule') or 'c1',mode='eval').body)!=dimensions[ligne['unite']]:raise ValueError('Unité du résultat incompatible avec le calcul.')
    return {**ligne,'quantite':format(resultat,'f')}

def dedoublonner(lignes):
    retenues={};conflits=[];bloquees=set()
    for l in lignes:
        cle=tuple(' '.join(l[k].casefold().split()) for k in ('lot','poste','niveau','local','unite'))
        if cle in bloquees:continue
        if cle in retenues and retenues[cle]['quantite']!=l['quantite']:
            conflits.append('Quantités contradictoires pour '+' / '.join(cle));retenues.pop(cle);bloquees.add(cle)
        else:retenues.setdefault(cle,l)
    return list(retenues.values()),conflits

import logging

MAX_RESERVES_RESUMEES=15
_NIVEAU=re.compile(r"(?<![A-Za-z0-9])(RDC|REZ[- ]DE[- ]CHAUSS[ÉE]E|R\s?\+\s?\d|SOUS[- ]SOL|SS\d?|COMBLES?|TOITURE)(?![A-Za-z0-9])",re.I)

def _cle_niveau(t):
    t=re.sub(r"\s+","",str(t or "").upper()).replace("REZ-DE-CHAUSSÉE","RDC").replace("REZ-DE-CHAUSSEE","RDC").replace("REZDECHAUSSÉE","RDC").replace("REZDECHAUSSEE","RDC")
    return t

def regrouper_reserves(reserves,maximum=MAX_RESERVES_RESUMEES):
    """(résumé, détail) : une remarque par IDÉE, la plus précise gardée, avec le nombre de fois où elle
    est revenue ; tout le reste part dans une feuille de détail — rien n'est perdu, rien n'est inventé."""
    from skills.documents_dossier import _mots_forts
    groupes=[]
    from collections import Counter
    fois=Counter(str(x).strip() for x in reserves if str(x).strip())
    for r in fois:
        mots=_mots_forts(r)
        for g in groupes:
            commun=len(mots&g['mots'])
            if mots and g['mots'] and commun/min(len(mots),len(g['mots']))>=.6:
                g['textes'].append(r);g['mots']|=mots;g['n']+=fois[r];break
        else:groupes.append({'mots':set(mots),'textes':[r],'n':fois[r]})
    groupes.sort(key=lambda g:-g['n'])
    resume=[]
    for g in groupes[:maximum]:
        texte=max(g['textes'],key=len)
        resume.append(texte+(' (remarque revenue '+str(g['n'])+' fois pendant la lecture)' if g['n']>1 else ''))
    if len(groupes)>maximum:resume.append(str(len(groupes)-maximum)+' autre(s) remarque(s) figurent dans la feuille « Réserves (détail) ».')
    detail=[t for g in groupes for t in g['textes']] if (len(groupes)>maximum or any(len(g['textes'])>1 for g in groupes)) else []
    return resume,detail

def couverture_des_niveaux(lignes,noms_des_pieces):
    """Ce que le classeur COUVRE, dit en tête : un métré du seul rez-de-chaussée présenté comme celui
    du bâtiment est le défaut le plus coûteux (relevé : 52 lignes, toutes « RDC », plans R+1 et R+2 fournis)."""
    couverts={_cle_niveau(l.get('niveau')) for l in lignes if str(l.get('niveau') or '').strip()}
    couverts={m.group(1) and _cle_niveau(m.group(1)) for c in couverts for m in [_NIVEAU.search(c)] if m}|{c for c in couverts if c in ('RDC',)}
    presents={_cle_niveau(m.group(1)) for nom in noms_des_pieces for m in _NIVEAU.finditer(str(nom))}
    manquants=sorted(presents-couverts)
    if not presents or not manquants:return ''
    return ('COUVERTURE INCOMPLÈTE : des lignes existent pour '+(', '.join(sorted(couverts)) or 'aucun niveau identifié')
            +' ; les pièces fournies concernent aussi '+', '.join(manquants)+', qui n’ont AUCUNE ligne. Ce classeur n’est pas le métré du bâtiment entier.')
logger=logging.getLogger('infra.documents')
PLAFOND_CONTEXTE = 12000
# Un nombre suivi d'une unité de métré : ce qu'une citation de quantitatif doit porter.
_QUANTITE = re.compile(r"\d[\d\s.,]*\s?(?:m²|m2|m³|m3|ml\b|m\.l|mL\b|m\b|cm\b|mm\b|u\b|U\b|unit[ée]s?\b|ens\b|forfait|ft\b|kg\b|t\b|l\b|%)", re.I)

_MOTS_DE_LOCALISATION=re.compile(r"(s[ée]jour|chambre|cuisine|salle|\bwc\b|sanitaire|pi[èe]ces?\s+(humides?|s[èe]ches?|principales?)|d[ée]gagement|entr[ée]e|cellier|logement|loggia|terrasse|balcon|hall|palier|carrel|fa[iï]ence|pvc|souple|lino|parquet|stratifi|moquette|\bu[234]s?\b|\bp[23]\b|localisation|rev[êe]tement)",re.I)
_ARTICLE=re.compile(r'(?m)^\s*(\d+(?:\.\d+){1,4})\.?\s+(\S[^\n]{3,140})$')
_LOCALISATION=re.compile(r'Localisation\s*:',re.I)
# Ce qui ne se mesure PAS par la surface au sol d'une pièce : ouvrages muraux ou ponctuels.
_NON_SURFACIQUE=re.compile(r"(fa[iï]ence|mural|plinthe|douche|paillasse|tablier|natte|forme de pente|siphon|baguette|joint|cr[ée]dence|tapis|seuil|nez de marche|escalier|podotactile|surbau)",re.I)
_EXCLUSION=re.compile(r"(hormis|sauf|except|à l[’']exception|hors)",re.I)
_EXTERIEUR=re.compile(r'(terrasse|loggia|balcon|jardin|patio|cour\b)',re.I)

def clauses_de_localisation(ref,texte):
    """LES CLAUSES « Localisation : … » D'UN CCTP, avec l'article qui les porte.
    Essai réel du 17/09 : servi en vrac, le modèle a posé du PVC dans les salles
    d'eau alors que le CCTP écrit « hormis chambres et pièces humides revêtues en
    carrelage », et du lissage à tous les niveaux pour « Sols des logements à
    RDC ». La clause est LA phrase qui dit où va un ouvrage : on la sert entière,
    rattachée à son article et à sa désignation."""
    clauses=[]
    for m in _LOCALISATION.finditer(texte):
        avant=texte[max(0,m.start()-2600):m.start()]
        articles=list(_ARTICLE.finditer(avant))
        article=(articles[-1].group(1)+' '+articles[-1].group(2).strip()) if articles else ''
        puce=max(avant.rfind('\n-\n'),avant.rfind('\n- \n'),avant.rfind('\n-\r\n'))
        designation=' '.join(avant[(puce+1 if puce>=0 else (articles[-1].end() if articles else max(0,len(avant)-300))):].split())
        designation=re.sub(r'\.{4,}',' … ',designation)[:320]
        suite=texte[m.start():m.start()+700]
        borne=re.search(r'\n\s*(?:\d+(?:\.\d+){1,4}\.?\s+\S|-\s*\n|=== Page)',suite[14:])
        clause=suite[:14+borne.start()] if borne else suite[:520]
        clauses.append({'preuve':ref,'article':article[:160],'designation':designation,'citation':clause.strip()[:620]})
    return clauses

async def _affecter_surfaces(uid,fil,tache,demande,surfaces,plans,a_metrer,lignes_tableaux,analyses,preuves,noms_sources):
    """LES SURFACES DES PIÈCES DEVIENNENT DES MÉTRÉS (17/09). Le DPGF du lot 12 du
    dossier réel ne porte AUCUNE quantité (13 postes « à métrer ») : c'est le
    travail demandé. Le code tient les surfaces exactes de chaque pièce ; il
    reste à dire QUEL poste couvre QUELS types de pièces, à QUELS niveaux — c'est
    la clause « Localisation » du CCTP qui le dit. Le modèle ne fait que cela,
    en désignant la clause qui le prouve (le serveur la recopie, rien n'est
    retapé). Les quantités, elles, ne passent jamais par lui.
    Rend (lignes, réserves, {(lot, poste) métrés}, lignes de contrôle)."""
    if not surfaces:return [],[],set(),[]
    from skills.tableaux_quantites import famille_de_piece
    noms_plans=' ; '.join('« '+p['nom']+' »' for p in plans)
    logements=len({x['logement'] for x in surfaces if x['type']})
    total=sum((x['surface'] for x in surfaces if x['type']),Decimal(0))
    reserves=['Surfaces des pièces relues par le code dans '+noms_plans+' : '+str(len(surfaces))+' pièce(s), '+str(logements)+' logement(s), '+format(total,'f').replace('.',',')
              +' m² habitables ; la somme de chaque logement est contrôlée sur le total écrit par l’architecte (feuille « Surfaces des pièces »). Ce sont des surfaces HABITABLES : ni déduction d’emprise, ni seuils, ni plinthes, ni remontées.']
    postes=[{**a,'_chiffre':None} for a in a_metrer if a['unite']=='m²']+[{'lot':l['lot'],'poste':l['poste'],'unite':'m²','local':l.get('local') or '','_chiffre':l['quantite']} for l in lignes_tableaux if l.get('unite')=='m²']
    postes=[a for a in postes if not _NON_SURFACIQUE.search(a['poste'])]
    if not postes:
        return [],reserves+['Aucun poste en m² dans les DPGF fournis : les surfaces des pièces sont livrées telles quelles, sans affectation à un poste.'],set(),[]
    familles={}
    for x in surfaces:
        f=familles.setdefault(famille_de_piece(x['piece']),{'pieces':0,'surface':Decimal(0),'niveaux':set()})
        f['pieces']+=1;f['surface']+=x['surface'];f['niveaux'].add(x['niveau'] or 'Non précisé')
    niveaux_connus=sorted({x['niveau'] or 'Non précisé' for x in surfaces})
    lots={m.group().lstrip('0') for a in postes for m in [re.search(r'\d+',a['lot'])] if m}
    def _du_lot(nom):
        m=re.search(r'lot\s*n?°?\s*0*(\d+)',str(nom or ''),re.I);return bool(m and m.group(1) in lots)
    passages=[]
    for ref,texte in preuves.items():
        if ':q' in ref or not _du_lot(noms_sources.get(ref.split(':')[0])):continue
        passages+=clauses_de_localisation(ref,texte)
    if not passages:
        # Un CCTP sans le mot « Localisation » : on se rabat sur les passages que la lecture a relevés.
        for a in analyses:
            for f in a['faits']:
                cit=str(f.get('citation') or '')
                if cit and a.get('preuve') in preuves and _du_lot(a.get('nom')) and _MOTS_DE_LOCALISATION.search(cit):
                    passages.append({'preuve':a['preuve'],'article':'','designation':str(f.get('fait') or '')[:200],'citation':cit[:620]})
    retenus=[];taille=0
    for x in passages:
        cout=len(x['citation'])+len(x['designation'])+len(x['article'])+60
        if taille+cout>26000:break
        retenus.append({'n':len(retenus),**x});taille+=cout
    if not retenus:
        return [],reserves+['Aucune clause du CCTP ne dit quel revêtement va dans quelle pièce : les surfaces sont livrées sans affectation à un poste. Ajoute le CCTP du lot ou le tableau de localisation.'],set(),[]
    cle='affectation_surfaces:v2'
    from skills.documents_dossier import dire
    await dire(uid,fil,tache,'rapprochement de '+str(len(postes))+' poste(s) du DPGF avec les pièces des logements, d’après '+str(len(retenus))+' clause(s) « Localisation » du CCTP')
    avis=await asyncio.to_thread(dossiers.etape,uid,fil,tache,cle)
    if not avis:
        def verifier(r):
            if not isinstance(r.get('affectations'),list) or not isinstance(r.get('sans_surface'),list):raise ValueError('affectations[] et sans_surface[] obligatoires.')
            for x in r['affectations']:
                if type(x.get('poste')) is not int or not 0<=x['poste']<len(postes):raise ValueError('Indice de poste invalide.')
                if type(x.get('clause')) is not int or not 0<=x['clause']<len(retenus):raise ValueError('Indice de clause invalide : désigne la clause « Localisation » qui PROUVE l’affectation.')
                if not isinstance(x.get('familles'),list) or not x['familles'] or any(f not in familles for f in x['familles']):raise ValueError('familles[] doit reprendre des noms EXACTS de la liste fournie : '+str(x.get('familles'))[:120])
                if not isinstance(x.get('niveaux'),list) or any(n not in niveaux_connus for n in x['niveaux']):raise ValueError('niveaux[] ne peut contenir que '+', '.join(niveaux_connus)+' (liste vide = tous les niveaux).')
        try:
            avis=await _json('Tu relies des POSTES de DPGF (en m²) aux TYPES DE PIÈCES des logements dont le code a relevé les surfaces. '
                'Pour CHAQUE poste : trouve SA clause « Localisation » (même article, même désignation), puis donne les familles de pièces qu’elle couvre et les niveaux. '
                'Applique la clause À LA LETTRE : ses exclusions (« hormis », « sauf », « excepté » : une pièce exclue n’entre pas), ses niveaux (« à RDC », « R+1 et R+2 », « uniquement »), ses pièces nommées (« SDB, SDE, WC et celliers »). '
                '« Logements » ou « sols des logements » = TOUTES les pièces intérieures des logements, aux niveaux dits. « Hall d’entrée », « circulations », « paliers » désignent les parties COMMUNES de l’immeuble, jamais l’entrée ni le dégagement d’un logement. '
                '« Pièces humides » = salles de bains, salles d’eau, WC (et celliers si la clause le dit). Une clause qui ne nomme aucune pièce de logement ne couvre rien. '
                'Le tableau ne porte QUE les pièces des logements et leurs annexes privatives : un poste localisé seulement dans les parties communes, halls, circulations, escaliers ou locaux techniques va dans sans_surface. '
                'Un poste peut avoir PLUSIEURS clauses (exemple : un même revêtement décrit en deux alinéas, l’un pour les chambres, l’autre pour le reste) : donne alors une affectation par clause. '
                'Une pièce ne porte qu’UN revêtement de sol : quand une clause d’un autre lot NOMME une pièce (« carrelage : SDB, SDE, WC et celliers »), cette pièce n’entre pas dans un revêtement défini par « toutes les pièces hormis… ». '
                'Une annexe EXTÉRIEURE (terrasse, loggia, balcon) n’est pas une « pièce du logement » : elle n’entre que si la clause la nomme. '
                'Si la clause couvre des logements ET d’autres lieux, affecte les pièces de logement et mets "partiel": true. Un ouvrage ponctuel (douche, paillasse, tablier, faïence murale) ne se mesure pas par la surface d’une pièce : sans_surface. '
                'Ne devine jamais d’après l’usage : sans clause, sans_surface. '
                'Schéma {"affectations":[{"poste":0,"clause":0,"familles":["noms exacts"],"niveaux":[],"partiel":false}],"sans_surface":[{"poste":0,"raison":"..."}]}. '
                '`poste` et `clause` sont les INDICES (n) des listes fournies ; niveaux vide = tous les niveaux ; ne retape aucun texte.',
                {'demande':demande,'niveaux':niveaux_connus,
                 'familles':[{'famille':k,'pieces':v['pieces'],'niveaux':sorted(v['niveaux']),**({'annexe_exterieure':True} if _EXTERIEUR.search(k) else {})} for k,v in sorted(familles.items())],
                 'postes':[{'n':i,'lot':a['lot'],'poste':a['poste'],'rubrique':a.get('local') or ''} for i,a in enumerate(postes)],
                 'clauses':[{'n':x['n'],'article':x['article'],'designation':x['designation'],'localisation':x['citation']} for x in retenus]},verifier)
            await asyncio.to_thread(dossiers.etape,uid,fil,tache,cle,avis)
        except Exception as e:
            logger.warning('Affectation des surfaces impossible (%s) : %s',type(e).__name__,str(e)[:200])
            return [],reserves+['L’affectation des pièces aux postes n’a pas pu être établie ('+type(e).__name__+') : les surfaces sont livrées dans leur feuille, sans ligne de métré.'],set(),[]
    # CE QU'UN AUTRE LOT NOMME NE REVIENT PAS PAR UNE CLAUSE GÉNÉRALE. Essai réel :
    # le sol souple va dans « toutes les pièces hormis chambres et pièces humides
    # revêtues en carrelage », et le lot carrelage NOMME « SDB, SDE, WC et
    # celliers » : les celliers restaient pourtant sous le PVC (40 m² en trop).
    def _nommee(famille,texte):
        return bool(re.search(r'(?<![\w])'+re.escape(famille.casefold())+r's?(?![\w])',texte.casefold()))
    nommees={}
    for x in avis['affectations']:
        clause=retenus[x['clause']];lot=postes[x['poste']]['lot']
        for f in x['familles']:
            if _nommee(f,clause['citation']):nommees.setdefault(f,set()).add(lot)
    # Plusieurs clauses peuvent nourrir UN poste ; une pièce n'y entre qu'une fois.
    par_poste={}
    for x in avis['affectations']:
        clause=retenus[x['clause']]
        if _EXCLUSION.search(clause['citation']):
            x={**x,'familles':[f for f in x['familles'] if not (nommees.get(f,set())-{postes[x['poste']]['lot']}) or _nommee(f,clause['citation'].split('hormis')[0])]}
        # GARDE MÉCANIQUE : une terrasse n'est pas une « pièce du logement ». Essai
        # réel : « toutes les pièces des logements hormis… » avait ramené loggias,
        # balcons et terrasses sous un sol PVC.
        familles_x=[f for f in x['familles'] if not _EXTERIEUR.search(f) or _EXTERIEUR.search(clause['citation'])]
        groupe=par_poste.setdefault(x['poste'],{'pieces':{}, 'perimetres':[], 'partiel':False})
        for piece in surfaces:
            if famille_de_piece(piece['piece']) in familles_x and (not x['niveaux'] or (piece['niveau'] or 'Non précisé') in x['niveaux']):
                groupe['pieces'].setdefault((piece['logement'],piece['piece'],str(piece['surface'])),(piece,clause))
        if familles_x:
            groupe['perimetres'].append(', '.join(familles_x)+(' — '+', '.join(x['niveaux']) if x['niveaux'] else ' — tous niveaux'))
            groupe['partiel']=groupe['partiel'] or bool(x.get('partiel'))
    lignes=[];metres=set();vus=set(par_poste);controle=[]
    for indice,groupe in par_poste.items():
        if not groupe['pieces']:continue
        poste=postes[indice];retenues=list(groupe['pieces'].values())
        somme=sum((piece['surface'] for piece,_ in retenues),Decimal(0))
        perimetre=' + '.join(dict.fromkeys(groupe['perimetres']))+(' ; logements SEULEMENT, la clause couvre aussi d’autres lieux' if groupe['partiel'] else '')
        if poste['_chiffre'] is not None:
            # Poste déjà chiffré par le maître d'œuvre : on ne le double pas, on le CONTRÔLE.
            controle.append([poste['lot'],poste['poste'],float(Decimal(poste['_chiffre'])),float(somme),float(somme-Decimal(poste['_chiffre'])),perimetre,' // '.join(dict.fromkeys(c['citation'] for _,c in retenues))[:1500]])
            continue
        metres.add((poste['lot'],poste['poste']))
        for piece,clause in retenues:
            lignes.append({'lot':poste['lot'],'poste':poste['poste'],'niveau':piece['niveau'] or 'Non précisé',
                           'local':'Logement '+piece['logement']+(' ('+piece['type']+')' if piece['type'] else '')+' — '+piece['piece'],
                           'unite':'m²','formule':'c1','quantite':format(piece['surface'],'f'),
                           'operandes':[{'valeur':format(piece['surface'],'f'),'unite':'m²','preuve':piece['preuve'],'citation':piece['citation']}],
                           'affectation':{'preuve':clause['preuve'],'citation':(clause['article']+' — ' if clause['article'] else '')+clause['citation']},
                           'lecture':'Surface habitable du tableau de l’architecte (lue par le code, somme contrôlée) ; pièces retenues d’après la clause « Localisation » citée'+(' — logements seulement, le poste couvre aussi d’autres lieux' if groupe['partiel'] else '')+' ; à contrôler sur le plan « nature des sols »','_tableau':True})
        reserves.append('« '+poste['poste']+' » : '+format(somme,'f').replace('.',',')+' m² relevés sur '+str(len(retenues))+' pièce(s) ('+perimetre+').')
    if lignes:reserves.append(str(len(lignes))+' ligne(s) de métré viennent des surfaces des pièces, retenues d’après les clauses « Localisation » du CCTP : contrôle-les sur les plans « NATURE DES SOLS » (une pièce peut porter deux revêtements) et ajoute chutes, seuils et plinthes.')
    if controle:reserves.append(str(len(controle))+' poste(s) déjà chiffré(s) par le DPGF sont comparés aux surfaces relevées dans la feuille « Contrôle du DPGF » (écart = relevé − DPGF).')
    for x in avis.get('sans_surface') or []:
        if isinstance(x,dict) and type(x.get('poste')) is int and 0<=x['poste']<len(postes) and postes[x['poste']]['_chiffre'] is None and x['poste'] not in vus:
            reserves.append('« '+postes[x['poste']]['poste']+' » reste à métrer sur plan : '+str(x.get('raison') or 'aucune clause ne le situe dans une pièce de logement')[:220])
    return lignes,reserves,metres,controle

async def _produire(data,user):
    uid,fil=dossiers.identite(getattr(user,'id',None),data.get('_fil'))
    demande=str(data.get('_demande_utilisateur') or data.get('demande') or '')
    ids=data.get('sources') or [s['id'] for s in dossiers.manifeste(uid,fil)]
    if not ids:raise ValueError('Ajoute les plans, CCTP et tableaux au dossier avant de demander le quantitatif.')
    sources=await asyncio.to_thread(dossiers.sources,uid,fil,ids)
    from skills.documents_dossier import verifier_acces
    await verifier_acces(sources,user)
    tache=data['_tache_quantitatif']
    from skills.documents_dossier import completer_visuels
    sources=await completer_visuels(uid,fil,tache,sources,user,demande)
    analyses=await _analyses(uid,fil,tache,demande,sources)
    # Les contraintes et affectations sont communes à tous les fragments à quantifier.
    # LE CONTEXTE COMMUN SE BORNE (17/09). Il portait TOUTES les analyses du
    # dossier, entières, dans CHAQUE appel : « 181 666 caractères à contrôler »
    # pour un fragment utile de 5 000. Chaque appel expirait, aucune étape
    # n'était jamais acquise, et la file rejouait le tout huit fois. On ne
    # transmet plus que les EXIGENCES (le fait, sa référence), d'abord celles de
    # la pièce du fragment puis du même lot, dans un plafond fixe.
    exigences=[{'source':a.get('source'),'preuve':a.get('preuve'),'fait':str(f.get('fait') or f.get('texte') or '')[:280]}
               for a in analyses for f in a['faits'] if f.get('nature')=='exigence']
    noms={s['id']:s['nom'] for s in sources}
    def _lot(nom):
        m=re.search(r'lot\s*n?°?\s*(\d+)',str(nom or ''),re.I);return m.group(1) if m else None
    def contexte_pour(source):
        lot=_lot(source['nom'])
        rang=lambda e:(0 if e['source']==source['id'] else 1 if lot and _lot(noms.get(e['source']))==lot else 2)
        retenus=[];taille=0
        for e in sorted(exigences,key=rang):
            cout=len(e['fait'])+60
            if taille+cout>PLAFOND_CONTEXTE:break
            retenus.append(e);taille+=cout
        return retenus
    fragments=[];preuves={};sans_quantite=[]
    # LES TABLEAUX CHIFFRÉS SE LISENT PAR LE CODE (17/09, décision de Noa : « les
    # Excel lus par le code, et en cas de problème seulement par l'IA »). Un DPGF
    # reconnu donne ses lignes exactes, sans jeton ; un classeur SANS en-tête
    # reconnaissable reste confié à la lecture habituelle, plus bas.
    from skills.tableaux_quantites import lire as lire_tableau
    lignes_tableaux=[];a_metrer=[];lus_par_code=set()
    for s in sources:
        tableau=lire_tableau(s)
        if tableau['entete']:
            lus_par_code.add(s['id']);lignes_tableaux+=tableau['lignes'];a_metrer+=tableau['a_metrer']
            logger.info('Tableau « %s » lu par le code : %d ligne(s) chiffrée(s), %d poste(s) à métrer',s['nom'][:60],len(tableau['lignes']),len(tableau['a_metrer']))
    # LE TABLEAU DES SURFACES D'UN PLAN SE LIT PAR LE CODE AUSSI (17/09). Le plan
    # « 12 SURFACES » du DCE réel porte la surface de chaque pièce des 29 logements,
    # colonne par colonne : aucun nombre n'y touche son unité (tout était rejeté),
    # et la somme de chaque logement retombe sur le total écrit par l'architecte.
    from skills.tableaux_quantites import lire_surfaces
    surfaces=[];plans_surfaces=[]
    for s in sources:
        if s['id'] in lus_par_code:continue
        releve=lire_surfaces(s)
        if releve['blocs']>=3:
            surfaces+=releve['pieces'];plans_surfaces.append(s);lus_par_code.add(s['id'])
            logger.info('Tableau des surfaces « %s » lu par le code : %d pièce(s), %d bloc(s) dont la somme est contrôlée',s['nom'][:60],len(releve['pieces']),releve['blocs'])
    # La lecture visuelle du MÊME tableau ne repasse pas par le modèle : elle
    # redonnerait les mêmes pièces sous d'autres libellés — un double compte.
    lus_par_code|={s['id'] for s in sources if s.get('_source_originale') in {p['id'] for p in plans_surfaces}}
    deja_releve=('Les surfaces des PIÈCES DES LOGEMENTS sont déjà relevées par le code depuis le tableau des surfaces de l’architecte : '
                 'ne rends AUCUNE ligne de surface de pièce de logement. Relève seulement ce que ce tableau ne porte pas : parties communes, halls, paliers, escaliers, locaux communs, longueurs, hauteurs, quantités comptées.') if surfaces else ''
    for s in sources:
        if s['id'] in lus_par_code:continue
        for f in dossiers.fragments(s['contenu'],taille=5000):
            cle=s['id']+':q'+str(f['numero']);texte=f['texte']
            preuves[cle]=texte
            # Une ligne de quantitatif exige une citation « portant valeur et
            # unité » : un fragment qui n'en contient aucune (règlement, clauses
            # administratives) ne peut rien rendre — inutile de payer l'appel.
            if _QUANTITE.search(texte):fragments.append((s,cle,texte))
            else:sans_quantite.append(cle)
        for f in dossiers.fragments(s['contenu']):preuves[f"{s['id']}:{f['numero']}"]=f['texte']
    await asyncio.to_thread(dossiers.etape,uid,fil,tache,'suivi_quantitatif',{'passages':len(fragments),'tableaux':len(lus_par_code),'lignes_tableaux':len(lignes_tableaux)+len(surfaces)})
    semaphore=asyncio.Semaphore(3)
    async def extraire(s,cle,texte):
        ancienne=await asyncio.to_thread(dossiers.etape,uid,fil,tache,cle)
        if ancienne:return ancienne
        async with semaphore:
            from skills.documents_dossier import dire
            await dire(uid,fil,tache,'relevé des quantités dans « '+s['nom'][:70]+' » (passage '+cle.rsplit(':q',1)[-1]+')')
            passes=[0]
            def verifier(r):
                if not isinstance(r.get('lignes'),list) or not isinstance(r.get('reserves'),list):raise ValueError('lignes[] et reserves[] obligatoires.')
                # UNE LIGNE FAUSSE NE FAIT PLUS PERDRE LES LIGNES JUSTES (17/09). La
                # première réponse invalide est renvoyée au modèle avec ses raisons
                # (il corrige souvent) ; à la seconde, on GARDE ce qui est prouvé et
                # l'on écarte le reste en le disant — avant, tout le passage était perdu.
                passes[0]+=1;bonnes=[];raisons=[]
                for l in r['lignes']:
                    try:bonnes.append(_verifier_ligne(l,preuves))
                    except (ValueError,TypeError,KeyError,ArithmeticError,SyntaxError) as e:
                        raisons.append('« '+str((l or {}).get('poste') if isinstance(l,dict) else l)[:80]+' » : '+str(e)[:160])
                if raisons and passes[0]==1:raise ValueError('Lignes à corriger : '+' ; '.join(raisons[:6]))
                r['lignes']=bonnes
                r['reserves']=list(r['reserves'])+['Ligne écartée faute de preuve suffisante — '+x for x in raisons[:8]]
            r=await _json('Établis les lignes de quantitatif prouvables à partir de CE fragment et des affectations du dossier. '
                'Ne réemploie pas les quantités d’un ancien exemple. Ne confonds pas surface habitable et surface de revêtement : '
                'sur un PLAN ou un tableau de surfaces, une surface de pièce lisible SE RELÈVE quand même (niveau = le niveau du plan, local = logement et pièce), '
                'avec la réserve « surface de pièce : déductions et relevés de plinthes non faits » ; son affectation peut citer le libellé de la pièce ou du logement tel qu’il est écrit sur le plan. '
                'niveau et local sont facultatifs quand la pièce ne les donne pas. '
                'Pas de faïence calculée sans hauteur/périmètre, pas de plinthes sans longueur justifiée. '
                'Un même local/poste doit avoir une clé stable entre plans et tableau. Donne toutes les lignes pertinentes du fragment ; '
                'réserves précises pour les postes non déterminables. Schéma {"lignes":[{"lot":"...","poste":"...","niveau":"...",'
                '"local":"...","unite":"m²","formule":"c1","operandes":[{"valeur":"42.5","unite":"m²","preuve":"référence",'
                '"citation":"texte exact portant valeur et unité"}],"affectation":{"preuve":"référence","citation":"texte exact du CCTP/localisation"}}],"reserves":["..."]}. '
                'La formule emploie UNIQUEMENT c1, c2, etc. dans l’ordre des opérandes et les signes + - * / : '
                'pour longueur fois largeur, écris c1*c2, jamais longueur x largeur. '
                'Les identifiants, dates et numéros de plan ne sont pas des quantités. Une estimation visuelle est une réserve, pas une mesure exacte.',
                {'demande':demande,'source':s['nom'],'preuve_fragment':cle,'fragment':texte,'affectations':contexte_pour(s),**({'deja_releve':deja_releve} if deja_releve else {})},verifier)
            await asyncio.to_thread(dossiers.etape,uid,fil,tache,cle,r);return r
    lus=await asyncio.gather(*(extraire(s,c,t) for s,c,t in fragments),return_exceptions=True)
    erreurs=['Un fragment n’a pas pu être quantifié ('+type(x).__name__+').' for x in lus if isinstance(x,BaseException)]
    bons=[x for x in lus if isinstance(x,dict)]
    lignes_surfaces,reserves_surfaces,metres,controle_dpgf=await _affecter_surfaces(uid,fil,tache,demande,surfaces,plans_surfaces,a_metrer,lignes_tableaux,analyses,preuves,noms)
    a_metrer=[a for a in a_metrer if (a['lot'],a['poste']) not in metres]
    lignes,conflits=dedoublonner(lignes_tableaux+lignes_surfaces+[l for r in bons for l in r['lignes']])
    # LES RÉSERVES DU MODÈLE SE REGROUPENT (17/09). Le quantitatif réel de La Teste, rejoué à blanc :
    # 102 réserves, la même idée (« aucune hauteur, pas de faïence calculable ») redite à chaque
    # fragment. Celles que le CODE écrit (couverture, dossier lu, contrôle du DPGF) restent entières.
    reserves_du_modele=[str(x) for r in bons for x in r['reserves']]
    reserves=list(dict.fromkeys(reserves_surfaces+conflits+erreurs))
    if a_metrer:
        reserves.append(str(len(a_metrer))+' poste(s) du DPGF n’ont PAS de quantité dans le tableau fourni et restent à métrer : '
                        +' ; '.join(a['poste']+' ('+a['unite']+')' for a in a_metrer[:25])+(' …' if len(a_metrer)>25 else ''))
    charge=data.get('_dossier_charge')
    if charge:
        if charge.get('introuvable'):reserves.append('Le dossier « '+charge['dossier']+' » n’a pas pu être ouvert sur le serveur ('+charge['introuvable']+') : seules les pièces jointes à la conversation ont été lues.')
        else:
            reserves.append('Dossier du serveur lu : « '+charge['dossier']+' » — '+str(charge.get('vus',0))+' fichier(s) vus, '+str(len(charge['ajoutes']))+' chargé(s)'
                            +(' pour les lots '+', '.join(str(n) for n in charge['lots'])+(' (déduits du métier de l’entreprise, la demande n’en nommait aucun)' if charge.get('lots_deduits') else '') if charge.get('lots') else ' (aucun lot précisé : précise-les pour cibler le relevé)')
                            +(' ; écartés : '+', '.join(str(v)+' '+k for k,v in charge['ecartes'].items()) if charge.get('ecartes') else '')+'.')
            if charge['ignores']:reserves.append('Fichiers du dossier non lus : '+' ; '.join(charge['ignores'][:20])+(' …' if len(charge['ignores'])>20 else ''))
    if sans_quantite:reserves.append(str(len(sans_quantite))+' passage(s) sans aucun nombre suivi d’une unité (clauses, règlement) n’ont pas été interrogés : ils ne peuvent porter aucune ligne de métré.')
    # Le contrôle final reçoit les exigences, bornées elles aussi.
    contexte_global=[];_t=0
    for e in exigences:
        if _t+len(e['fait'])+60>PLAFOND_CONTEXTE*2:break
        contexte_global.append(e);_t+=len(e['fait'])+60
    visuels={s['id'] for s in sources if 'lecture visuelle' in s['nom'].casefold() or 'analyse visuelle' in s['contenu'][:150].casefold()}
    for l in lignes:
        if l.get('_tableau'):continue
        l['lecture']='Visuelle : à contrôler sur le plan' if any(o['preuve'].split(':')[0] in visuels for o in l['operandes']) else 'Texte extrait'
    if any(l['lecture'].startswith('Visuelle') for l in lignes):reserves.append('Les quantités issues de lectures visuelles doivent être contrôlées sur les plans originaux ; la citation vérifie la transcription conservée, pas la mesure physique.')
    a_controler=[l for l in lignes if not l.get('_tableau')]
    if a_controler:
        lignes_du_code=[l for l in lignes if l.get('_tableau')];lignes=a_controler
        avis=await _json('Vérifie chaque ligne quantitative : affectation réelle au poste et au local, absence de double comptage, ancienne opération non réemployée, cotes et unités lisibles. '
            'Pour chaque ligne invalide, donne son indice zéro-based dans rejeter et une réserve précise. JSON {"rejeter":[0],"reserves":["..."]}. Une simple surface habitable ne prouve pas automatiquement une surface de revêtement.',
            {'demande':demande,'lignes':lignes,'affectations':contexte_global})
        rejeter=avis.get('rejeter');notes=avis.get('reserves')
        if not isinstance(rejeter,list) or any(type(i) is not int or not 0<=i<len(lignes) for i in rejeter) or not isinstance(notes,list):raise ValueError('Contrôle quantitatif invalide.')
        lignes=lignes_du_code+[l for i,l in enumerate(lignes) if i not in rejeter];reserves_du_modele.extend(str(n) for n in notes)
    resume,detail_reserves=regrouper_reserves(reserves_du_modele)
    couverture=couverture_des_niveaux(lignes,[s['nom'] for s in sources])
    reserves=([couverture] if couverture else [])+reserves+resume
    if not lignes:return {'ok':True,'outcome':'partial','production_verifiee':False,'definitif':not erreurs,'reserves':reserves or ['Aucune quantité avec affectation et unité suffisamment prouvées.'],'a_faire':'Identifie les pages et cotes manquantes ; analyser_plan_source permet de lire le dessin. Ne livre pas un inventaire de fichiers comme métré.'}
    from bureautique import atelier
    from skills.bureau import terminer_document
    titres=['Lot','Poste','Niveau','Local','Quantité','Unité','Formule','Sources et citations','Lecture']
    rows=[[l[k] for k in ('lot','poste','niveau','local','quantite','unite','formule')]+[json.dumps({'operandes':l['operandes'],'affectation':l['affectation']},ensure_ascii=False),l['lecture']] for l in lignes]
    for l in lignes:l.pop('_tableau',None)
    for row in rows:row[4]=float(Decimal(row[4]))
    totaux={}
    for l in lignes:
        k=(l['lot'],l['poste'],l['niveau'],l['unite']);totaux[k]=totaux.get(k,Decimal(0))+Decimal(l['quantite'])
    feuille_surfaces=[{'bloc':'feuille','nom':'Surfaces des pièces','colonnes_numeriques':[4],'entetes':['Niveau','Logement','Type','Pièce','Surface (m²)','Contrôle'],
                       'lignes':[[x['niveau'] or 'Non précisé',x['logement'],x['type'],x['piece'],float(x['surface']),'Somme du logement = total écrit sur le plan'] for x in surfaces]}] if surfaces else []
    blocs=[{'bloc':'feuille','nom':'Détail','colonnes_numeriques':[4],'entetes':titres,'lignes':rows},
           {'bloc':'feuille','nom':'Synthèse','colonnes_numeriques':[4],'entetes':['Lot','Poste','Niveau','Unité','Quantité'],'lignes':[[*k,float(v)] for k,v in sorted(totaux.items())]},
           *feuille_surfaces,
           *([{'bloc':'feuille','nom':'Contrôle du DPGF','colonnes_numeriques':[2,3,4],'entetes':['Lot','Poste','Quantité du DPGF (m²)','Surface relevée (m²)','Écart (m²)','Pièces retenues','Clause du CCTP'],'lignes':controle_dpgf}] if controle_dpgf else []),
           {'bloc':'feuille','nom':'Réserves','entetes':['Point à vérifier'],'lignes':[[r] for r in reserves] or [['Aucune réserve détectée ; contrôle métier humain requis.']]},
           *([{'bloc':'feuille','nom':'Réserves (détail)','entetes':['Remarque relevée pendant la lecture'],'lignes':[[r] for r in detail_reserves[:2000]]}] if detail_reserves else [])]
    preuves_lignes=[]
    for i,l in enumerate(lignes,1):
        for o in l['operandes']:
            for debut in range(0,len(o['citation']),1900):preuves_lignes.append([i,'Opérande',o['preuve'],o['valeur'],o['unite'],o['citation'][debut:debut+1900]])
        a=l['affectation']
        for debut in range(0,len(a['citation']),1900):preuves_lignes.append([i,'Affectation',a['preuve'],'','',a['citation'][debut:debut+1900]])
    for debut in range(0,len(preuves_lignes),5000):
        blocs.append({'bloc':'feuille','nom':'Preuves '+str(debut//5000+1),'entetes':['Ligne détail','Type','Source et fragment','Valeur','Unité','Citation'],'lignes':preuves_lignes[debut:debut+5000]})
    jeton=await asyncio.to_thread(dossiers.etape,uid,fil,tache,'jeton')
    if not jeton:
        jeton=await asyncio.to_thread(atelier.ouvrir,{'titre':data.get('titre') or 'Quantitatif sourcé','format':'xlsx'},uid,fil)
        await asyncio.to_thread(dossiers.etape,uid,fil,tache,'jeton',jeton)
    from bureautique.modele import normaliser_element
    attendus=[normaliser_element(b) for b in blocs];presents=list(atelier.elements(jeton))
    if presents!=attendus[:len(presents)]:
        await asyncio.to_thread(atelier.abandonner,jeton,uid)
        await asyncio.to_thread(dossiers.effacer_etapes,uid,fil,tache,['jeton'])
        raise ValueError('Le quantitatif a évolué ; nouveau rendu à la prochaine reprise.')
    if len(presents)<len(attendus):await asyncio.to_thread(atelier.ajouter,jeton,attendus[len(presents):],uid,False)
    r=await terminer_document({'document_id':jeton,'_fil':fil},user)
    from openpyxl import load_workbook
    wb=load_workbook(atelier.chemin_fichier(jeton,uid),read_only=True,data_only=True)
    try:
        if wb['Détail'].max_row!=len(lignes)+1:raise ValueError('Le classeur rendu ne contient pas toutes les lignes calculées.')
        if any(type(row[4]) not in (int,float) for row in wb['Détail'].iter_rows(min_row=2,values_only=True)):raise ValueError('Les quantités Excel ne sont pas numériques.')
    finally:wb.close()
    return {**r,'ok':True,'production_verifiee':True,'outcome':'partial' if reserves else 'success','lignes':len(lignes),
            'reserves':reserves,'sources_lues':len(sources),'fragments_quantifies':len(bons),'fragments_total':len(fragments),'a_faire':'Présente le quantitatif, sa couverture et les réserves. Il contient les quantités établies avec leurs preuves et calculs ; ne le prétends pas exhaustif si des postes restent en réserve.'}

async def produire_immediat(data,user):
    uid,fil=dossiers.identite(getattr(user,'id',None),data.get('_fil'))
    data=await asyncio.to_thread(dossiers.normaliser_selection,uid,fil,data)
    demande=str(data.get('_demande_utilisateur') or data.get('demande') or '')
    # LE CONTRAT SE FIGE AU PREMIER ESSAI (17/09). Un nouvel essai de la file porte
    # la `tache` du premier : on reprend SES pièces et SA demande. Sans cela, les
    # lectures visuelles ajoutées entre-temps changeaient la liste des pièces,
    # donc l'identité de la tâche — et toutes les étapes acquises étaient perdues.
    tache=data.get('tache')
    contrat=await asyncio.to_thread(dossiers.etape,uid,fil,tache,'contrat') if tache else None
    if contrat:
        demande=contrat['demande'];ids=contrat['sources'];charge=contrat.get('dossier')
    else:
        # LE DOSSIER NOMMÉ PAR LA DEMANDE EST OUVERT (17/09) : « les métrés du projet
        # "construction 29 lgts…" » ne lisait que les pièces jointes au chat.
        from skills.documents_dossier import dossier_du_travail
        charge=await dossier_du_travail(uid,fil,{**data,'tache':None},user,demande,'quantitatif')
        ids=data.get('sources') or []
        # Un dossier NOMMÉ et ouvert commande TOUTES les pièces du travail.
        if (charge and not charge.get('introuvable')) or not ids:ids=[s['id'] for s in dossiers.manifeste(uid,fil)]
        # SANS AUCUNE PIÈCE, LE REFUS EST DÉFINITIF ET DIT POURQUOI (18/09) : dossier ambigu ou
        # introuvable → la conversation reçoit la raison (les dossiers candidats), pas huit essais.
        if not ids:
            from skills.documents_dossier import _refus_definitif
            raise _refus_definitif('Aucune pièce pour ce quantitatif'+(' : « '+str(charge.get('dossier'))+' » n’a pas pu être ouvert ('+str(charge.get('introuvable') or 'aucun fichier lisible')+')' if charge else '')+'. Nomme le dossier exact du serveur, ou ajoute les pièces avec ajouter_source_dossier.')
        tache=hashlib.sha256(json.dumps(['quantitatif',demande,ids,data.get('titre')],ensure_ascii=False).encode()).hexdigest()[:24]
    from ressources.documents_file import associer_tache
    await asyncio.to_thread(associer_tache,uid,fil,tache)
    from stockage.verrous import verrou_fichier
    from ressources.registre import _chemin
    from bureautique import atelier
    with verrou_fichier(str(_chemin().parent),'document:'+uid+':'+tache,bloquant=False) as acquis:
        if not acquis:return {'ok':True,'en_cours':True,'production_verifiee':False,'tache':tache}
        fini=await asyncio.to_thread(dossiers.etape,uid,fil,tache,'livraison')
        if fini:
            if not atelier.chemin_fichier(fini['document_id'],uid):raise ValueError('Le quantitatif sauvegardé est devenu indisponible.')
            return fini
        if not contrat:await asyncio.to_thread(dossiers.etape,uid,fil,tache,'contrat',{'demande':demande,'sources':ids,'dossier':charge})
        try:r=await _produire({**data,'sources':ids,'_tache_quantitatif':tache,'_dossier_charge':charge},user)
        except Exception as e:
            # La file doit mesurer les étapes de ce quantitatif même après échec.
            e.tache_documentaire=tache
            raise
        r['tache']=tache
        if r.get('production_verifiee'):await asyncio.to_thread(dossiers.etape,uid,fil,tache,'livraison',r)
        return r

async def produire(data,user):
    from ressources.documents_file import soumettre
    uid,fil=dossiers.identite(getattr(user,'id',None),data.get('_fil'))
    return await asyncio.to_thread(soumettre,uid,fil,'quantitatif',data)

SKILLS={'produire_quantitatif' :Declaration(produire,'Produire un Excel de quantités depuis les plans, CCTP/DPGF et tableaux du dossier : détail par poste/local/niveau, unités et calculs vérifiés, synthèse et réserves. Aucun chiffre inventé. `dossier` : le NOM ou le chemin du dossier du serveur nommé par la demande (« les métrés du projet X », « à partir du dossier X ») — tous ses fichiers lisibles sont chargés, sous-dossiers compris, SANS les ajouter un par un. Les tableaux chiffrés (DPGF, DQE, BPU) sont lus par le code, ligne à ligne, avec leur cellule pour preuve. Ajouter les pièces avec ajouter_source_dossier et lire les plans graphiques avec analyser_plan_source avant si nécessaire.',requis=['demande'],optionnels=['sources','dossier','titre'],effet='ecriture_interne',expert='agent2',libelle='je calcule le quantitatif et vérifie ses sources')}
